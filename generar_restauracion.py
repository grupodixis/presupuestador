import os
import subprocess
import sys

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import openpyxl
except ImportError:
    install('openpyxl')
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

work_dir = r"C:\Users\pablo\Documents\Presupuestador"
os.makedirs(work_dir, exist_ok=True)

# ---------------------------------------------------------
# 1. GENERATE EXCEL
# ---------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "Presupuesto_Restauracion"

headers = ['Código', 'Descripción de la Unidad de Obra', 'Unidad', 'Cantidad', 'Precio Unitario (€)', 'Total (€)']
ws.append(headers)

header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
header_font = Font(bold=True)
for col in range(1, 7):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

data = [
    ['01.01', 'Implantación SS, Plan de Seguridad y Recursos Preventivos', 'Ud', 1, 8500.00],
    ['01.02', 'Montaje andamiaje multidireccional adaptado y medios verticales (x2 antenas)', 'Ud', 1, 65000.00],
    ['02.01', 'Protección perimetral ligera (Eliminado encapsulado pesado al usar tecnología láser)', 'Ud', 1, 8500.00],
    ['02.02', 'Decapado y saneamiento por tecnología LÁSER de alta potencia (cero residuos abrasivos)', 'm2', 4392.44, 55.00],
    ['03.01', 'Sistema anticorrosivo C5-M: Epoxi rica en zinc + Intermedia + Poliuretano', 'm2', 4392.44, 46.20],
    ['04.01', 'Sustitución de tornillería y perfiles LPN/IPN dañados (Estimado 8% masa)', 'kg', 8788.51, 8.50],
    ['05.01', 'Gestión ambiental de residuos: Filtros HEPA de limpieza láser y pequeña chatarra', 'Ud', 1, 2500.00]
]

for row_data in data:
    total = row_data[3] * row_data[4]
    row_data.append(total)
    ws.append(row_data)

ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 95
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 20
ws.column_dimensions['F'].width = 20

for row in range(2, len(data) + 2):
    ws.cell(row=row, column=5).number_format = '#,##0.00 €'
    ws.cell(row=row, column=6).number_format = '#,##0.00 €'

max_row = len(data) + 1
ws.cell(row=max_row+2, column=5, value="TOTAL PRESUPUESTO (€)").font = Font(bold=True)
total_cell = ws.cell(row=max_row+2, column=6, value=f"=SUM(F2:F{max_row})")
total_cell.font = Font(bold=True)
total_cell.number_format = '#,##0.00 €'

excel_path = os.path.join(work_dir, "Presupuesto_Restauracion_Antenas_Laser.xlsx")
wb.save(excel_path)
print(f"Excel guardado en {excel_path}")

# ---------------------------------------------------------
# 2. GENERATE HTML
# ---------------------------------------------------------
html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <title>Anteproyecto Restauración Antenas</title>
    <style>
        @page {{ margin: 25mm 20mm; }}
        body {{ font-family: 'Cambria', 'Georgia', serif; line-height: 1.6; margin: 0; color: #2b2b2b; font-size: 11pt; box-sizing: border-box; }}
        h1, h2, h3, h4 {{ font-family: 'Segoe UI', 'Helvetica Neue', sans-serif; color: #1a365d; page-break-after: avoid; }}
        .header {{ border-bottom: 2px solid #1a365d; padding-bottom: 15px; margin-bottom: 40px; display: table; width: 100%; }}
        .header-logo {{ display: table-cell; font-family: 'Segoe UI', sans-serif; font-size: 20pt; font-weight: 700; color: #1a365d; vertical-align: bottom; }}
        .header-logo span {{ color: #d9534f; }}
        .header-info {{ display: table-cell; text-align: right; font-size: 9pt; color: #555; vertical-align: bottom; font-family: 'Segoe UI', sans-serif; line-height: 1.4; }}
        .title-block {{ text-align: center; margin-bottom: 40px; }}
        .title-block h1 {{ font-size: 22pt; text-transform: uppercase; letter-spacing: 1px; margin-bottom: 10px; }}
        .title-block .subtitle {{ font-size: 14pt; color: #4a5568; font-family: 'Segoe UI', sans-serif; }}
        .section {{ margin-bottom: 25px; }}
        .section h2 {{ font-size: 14pt; border-bottom: 1px solid #cbd5e0; padding-bottom: 5px; margin-top: 25px; }}
        p, li {{ text-align: justify; }}
        ol, ul {{ padding-left: 20px; margin-top: 5px; margin-bottom: 5px; }}
        li {{ margin-bottom: 6px; }}
        .critical-note {{ background-color: #f7fafc; border-left: 4px solid #1a365d; padding: 15px 20px; margin: 15px 0; font-size: 10.5pt; }}
        .critical-note strong {{ color: #1a365d; font-family: 'Segoe UI', sans-serif; display: block; margin-bottom: 5px; font-size: 11pt; text-transform: uppercase; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 20px; font-family: 'Segoe UI', sans-serif; font-size: 9.5pt; page-break-inside: avoid; }}
        th, td {{ border: 1px solid #a0aec0; padding: 8px 10px; text-align: left; }}
        th {{ background-color: #edf2f7; color: #1a365d; font-weight: 600; text-transform: uppercase; }}
        .total-row {{ font-weight: bold; background-color: #e2e8f0; font-size: 11pt; }}
        .currency {{ text-align: right; white-space: nowrap; }}
        .footer {{ margin-top: 40px; text-align: center; font-size: 8pt; color: #718096; border-top: 1px solid #e2e8f0; padding-top: 15px; font-family: 'Segoe UI', sans-serif; }}
        .page-break {{ page-break-before: always; }}
    </style>
</head>
<body>
    <div class="header">
        <div class="header-logo">DIXIS<span>.</span> INGENIERÍA Y OBRAS</div>
        <div class="header-info">
            <strong>Ref. Doc:</strong> AP-2026-085-V3<br>
            <strong>Fecha:</strong> 20 de Agosto de 2026<br>
            <strong>Proyecto:</strong> Restauración de 2 Antenas Troposféricas<br>
            <strong>Ubicación:</strong> S'Enclusa, Menorca
        </div>
    </div>

    <div class="title-block">
        <h1>Proyecto Técnico y Valorado</h1>
        <div class="subtitle">Tratamiento Anticorrosivo Integral con Tecnología Láser</div>
    </div>
    
    <div class="section">
        <h2>1. Alcance y Objeto del Proyecto</h2>
        <p>El presente documento define la intervención técnica para la restauración integral de las dos estructuras parabólicas (antenas troposféricas) ubicadas en la antigua base militar. El objetivo es detener la degradación por corrosión ambiental y consolidar ambas estructuras utilizando tecnologías limpias.</p>
        <div class="critical-note">
            <strong>Dimensionamiento Exacto de las 2 Estructuras</strong>
            Se ha cuantificado el volumen total de intervención para las <strong>dos antenas gemelas</strong>: <strong>109.856,38 Kg</strong> de masa metálica con un desarrollo superficial colosal de <strong>4.392,44 m²</strong>.
        </div>
    </div>

    <div class="section">
        <h2>2. Metodología de Intervención (Sistema C5-M y Láser)</h2>
        <ol>
            <li><strong>Implantación y Medios Auxiliares:</strong> Montaje secuencial de andamiaje multidireccional adaptado a las cerchas traseras y uso de trabajos verticales (IRATA) para los discos reflectantes y soportes LNB.</li>
            <li><strong>Decapado y Limpieza LÁSER (Alta Potencia):</strong> Sustitución del tradicional chorreado abrasivo por <strong>decapado mediante tecnología láser de estado sólido</strong>. Esta técnica de vanguardia vaporiza la pintura deteriorada y el óxido estructural sin contacto ni impacto cinético, alcanzando una limpieza óptima del acero base sin generar residuos secundarios (arena contaminada). Los vapores son aspirados en el mismo cabezal mediante filtros HEPA.</li>
            <li><strong>Tratamiento Anticorrosivo Multicapa:</strong> Aplicación de un esquema de pintura diseñado para ambientes de alta salinidad extrema (Categoría C5-M): imprimación epoxídica rica en zinc, capa intermedia de epoxi micáceo y acabado en esmalte de poliuretano alifático.</li>
            <li><strong>Consolidación Mecánica:</strong> Sustitución de tornillería estructural fatigada por elementos nuevos de calidad 8.8 o superior, y sustitución de un estimado del <strong>8% de la perfilería total</strong> (mayormente cruces y entramados LPN80) que presenten una pérdida de sección estructural inasumible por la corrosión prolongada.</li>
        </ol>
    </div>

    <div class="section">
        <h2>3. Gestión Medioambiental (Reserva de la Biosfera)</h2>
        <p>La adopción de la <strong>tecnología láser</strong> representa un salto cualitativo en materia medioambiental, algo crítico al encontrarnos en una Reserva de la Biosfera. Se elimina por completo la necesidad de instalar encapsulados pesados de contención y se suprimen las toneladas de residuos de arena contaminada. La gestión de residuos se limitará a la entrega controlada de los filtros HEPA saturados (de los equipos de aspiración del láser) y la chatarra derivada de las sustituciones estructurales.</p>
    </div>

    <div class="section page-break">
        <h2>4. Valoración Económica</h2>
        <p>Valoración basada en la medición estructural exacta de casi 4.400 metros cuadrados de acero, considerando el ahorro en logística de contención y el incremento tecnológico del equipo láser portátil de alta potencia.</p>
        <table>
            <thead>
                <tr>
                    <th>Código</th>
                    <th>Descripción de la Unidad de Obra</th>
                    <th>Ud.</th>
                    <th>Cant.</th>
                    <th class="currency">Precio Unit.</th>
                    <th class="currency">Importe</th>
                </tr>
            </thead>
            <tbody>
"""

total_budget = 0
for row_data in data:
    code, desc, unit, cant, price, total = row_data
    total_budget += total
    html_content += f"""
                <tr>
                    <td>{code}</td>
                    <td>{desc}</td>
                    <td>{unit}</td>
                    <td>{cant:,.2f}</td>
                    <td class="currency">{price:,.2f} €</td>
                    <td class="currency">{total:,.2f} €</td>
                </tr>"""

html_content += f"""
                <tr class="total-row">
                    <td colspan="5" style="text-align: right; border-right: none;">PRESUPUESTO DE EJECUCIÓN MATERIAL (Sin IVA):</td>
                    <td class="currency">{total_budget:,.2f} €</td>
                </tr>
            </tbody>
        </table>
    </div>
    
    <div class="footer">
        Documento generado por DIXIS INGENIERÍA Y OBRAS - Departamento de Proyectos
    </div>
</body>
</html>
"""

html_path = os.path.join(work_dir, "Restauracion_Antenas_Laser.html")
pdf_path = os.path.join(work_dir, "Proyecto_Restauracion_Antenas_Laser.pdf")

with open(html_path, "w", encoding="utf-8") as f:
    f.write(html_content)

# ---------------------------------------------------------
# 3. GENERATE PDF
# ---------------------------------------------------------
edge_paths = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe"
]

browser_exe = None
for p in edge_paths:
    if os.path.exists(p):
        browser_exe = p
        break

if browser_exe:
    result = subprocess.run([
        browser_exe,
        "--headless",
        "--disable-gpu",
        "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}",
        html_path
    ], capture_output=True, text=True)
    if result.returncode == 0:
        print(f"PDF guardado en {pdf_path}")
    else:
        print(f"Error PDF: {result.stderr}")
else:
    print("Browser not found.")
