import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.sync_api import sync_playwright

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "P-2026-0047 Hotel Sur"

    # Estilos y Paleta Corporativa HAM
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=11, bold=True, color="B45309")
    regular_font = Font(name="Segoe UI", size=9.5)
    bold_font = Font(name="Segoe UI", size=9.5, bold=True)
    accent_total_font = Font(name="Segoe UI", size=11, bold=True, color="B45309")
    section_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Cabecera
    ws["A1"] = "HAM ESTRUCTURAS METÁLICAS MENORCA"
    ws["A1"].font = title_font
    ws["A2"] = "PRESUPUESTO TÉCNICO Y ECONÓMICO: P-2026-0047"
    ws["A2"].font = subtitle_font

    ws["A4"] = "Cliente:"
    ws["B4"] = "Hotel Sur"
    ws["A5"] = "Ubicación:"
    ws["B5"] = "Instalaciones Hotel Sur, Menorca"
    ws["A6"] = "Proyecto / Concepto:"
    ws["B6"] = "6 Partidas Hotel Sur 2026/2027: Vidrio, Inox 316L, Carpintería Aluminio, Separadores, Bastidores y Toldos"
    ws["A7"] = "Fecha y Validez:"
    ws["B7"] = "29/10/2026 · Validez: 30 días naturales"
    ws["A8"] = "Período Previsto:"
    ws["B8"] = "Temporada baja hotelera 2026 / 2027"

    for r in range(4, 9):
        ws[f"A{r}"].font = bold_font
        ws[f"B{r}"].font = regular_font

    # Resumen Parámetros Técnicos
    ws["A10"] = "ALCANCE GENERAL"
    ws["A10"].font = bold_font
    ws["A10"].fill = section_fill
    ws["B10"] = "6 partidas completas de renovación. Incluye desmontaje y retirada de elementos antiguos en partidas 2, 3, 4 y 6. Materiales de alta resistencia al ambiente marino (Inox AISI 316L, aluminio Qualicoat y vidrio de seguridad)."
    ws["B10"].font = regular_font
    ws["B10"].fill = section_fill

    # Tabla de partidas
    headers = ["Partida", "Concepto y Descripción Técnica", "Unidad", "Cantidad", "Precio Unitario (€)", "Importe (€)"]
    start_row = 12

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx not in [2] else "left", vertical="center")

    data = [
        ["01", "2 Fijos de cristal con perfil de mampara. Dimensiones: 2,64 x 2,64 m cada uno (Superficie total: 13,94 m²). Incluye perfilería de aluminio lacado/anodizado, vidrio de seguridad con cantos pulidos y sellado perimetral estanco EPDM.", "m2", 13.9392, 400.00],
        ["02", "Barandilla de cristal con bases en tubo 40x40x2 mm Acero Inoxidable AISI 316L (55 ML) con cristal laminar de seguridad 4+4 mm. Anclajes inox A4. Incluye retirada de la vieja barandilla atornillada a pared lateral y transporte a vertedero.", "ml", 55.00, 455.00],
        ["03", "Puerta de aluminio blanco calidad Qualicoat con cristal y fijo superior. Medidas hueco: altura 2,40 m y ancho 1,00 m. Incluye herrajes de uso intensivo, cerradura de seguridad y retirada de la vieja puerta.", "ud", 1.00, 2300.00],
        ["04", "Separadores de aluminio y panel blanco (3 uds). Medidas unitarias: 3,00 m ancho x 1,50 m alto (Superficie total: 13,50 m²). Incluye estructura perimetral, fijaciones ocultas y retirada de antiguo separador.", "m2", 13.50, 330.00],
        ["05", "2 Bastidores metálicos de perfil 80x40x2 mm forrados de madera (tablero) de 2,70 x 0,50 m amurados a pared mediante anclajes químicos. 1 ud para mesa o estante en comedor y 1 ud para expositor de tablets en recepción.", "ud", 2.00, 1250.00],
        ["06", "Suministro e instalación de 3 unidades de Toldo lateral retráctil de 3 m aprox., tejido técnico resinado (color a definir). Cofre de aluminio lacado con sistema de retroceso amortiguado. Incluye retirada de antiguos toldos.", "ud", 3.00, 500.00]
    ]

    for i, row in enumerate(data, start=start_row + 1):
        ws.cell(row=i, column=1, value=row[0]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=2, value=row[1]).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=i, column=3, value=row[2]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=4, value=row[3]).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=i, column=5, value=row[4]).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=i, column=6, value=f"=D{i}*E{i}").alignment = Alignment(horizontal="right", vertical="top")

        for col_idx in range(1, 7):
            c = ws.cell(row=i, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx in [4]:
                c.number_format = '#,##0.00'
            elif col_idx in [5, 6]:
                c.number_format = '#,##0.00 €'

    last_data_row = start_row + len(data)

    # Totales
    r_bi = last_data_row + 2
    ws.cell(row=r_bi, column=5, value="Base Imponible:").font = bold_font
    ws.cell(row=r_bi, column=6, value=f"=SUM(F{start_row+1}:F{last_data_row})").font = bold_font
    ws.cell(row=r_bi, column=6).number_format = '#,##0.00 €'

    r_iva = r_bi + 1
    ws.cell(row=r_iva, column=5, value="IVA (21%):").font = regular_font
    ws.cell(row=r_iva, column=6, value=f"=F{r_bi}*0.21").font = regular_font
    ws.cell(row=r_iva, column=6).number_format = '#,##0.00 €'

    r_tot = r_iva + 1
    ws.cell(row=r_tot, column=5, value="TOTAL PRESUPUESTO:").font = accent_total_font
    ws.cell(row=r_tot, column=6, value=f"=F{r_bi}+F{r_iva}").font = accent_total_font
    ws.cell(row=r_tot, column=6).number_format = '#,##0.00 €'

    # Ajuste de anchos de columna
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 65
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 18

    excel_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "presupuesto.xlsx"))
    wb.save(excel_path)
    print(f"Excel generado con exito: {excel_path}")
    return excel_path

def convert_html_to_pdf():
    base_dir = os.path.dirname(__file__)
    cliente_html = os.path.abspath(os.path.join(base_dir, "presupuesto-cliente.html"))
    cliente_pdf = os.path.abspath(os.path.join(base_dir, "presupuesto-cliente.pdf"))
    tecnico_html = os.path.abspath(os.path.join(base_dir, "presupuesto-interno-tecnico.html"))
    tecnico_pdf = os.path.abspath(os.path.join(base_dir, "presupuesto-interno-tecnico.pdf"))

    with sync_playwright() as p:
        browser = p.chromium.launch()

        # PDF Cliente
        page = browser.new_page()
        page.goto(f"file:///{cliente_html.replace(os.sep, '/')}", wait_until="networkidle")
        page.pdf(
            path=cliente_pdf,
            format="A4",
            print_background=True,
            margin={"top": "0mm", "bottom": "0mm", "left": "0mm", "right": "0mm"}
        )
        print(f"PDF Cliente generado: {cliente_pdf}")

        # PDF Tecnico
        page2 = browser.new_page()
        page2.goto(f"file:///{tecnico_html.replace(os.sep, '/')}", wait_until="networkidle")
        page2.pdf(
            path=tecnico_pdf,
            format="A4",
            print_background=True,
            margin={"top": "0mm", "bottom": "0mm", "left": "0mm", "right": "0mm"}
        )
        print(f"PDF Tecnico generado: {tecnico_pdf}")

        browser.close()

if __name__ == "__main__":
    create_excel()
    convert_html_to_pdf()
