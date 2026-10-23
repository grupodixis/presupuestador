import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.sync_api import sync_playwright

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "P-2026-0042 Escalera Galv"

    # Paleta de colores corporativa técnica (Azul pizarra / Slate / Cyan oscuro)
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=11, bold=True, color="0284C7")
    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    accent_total_font = Font(name="Segoe UI", size=12, bold=True, color="0284C7")
    section_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    double_bottom_border = Border(
        top=Side(style='thin', color='0F172A'),
        bottom=Side(style='double', color='0F172A')
    )

    # Cabecera Presupuesto
    ws["A1"] = "HAM ESTRUCTURAS METÁLICAS"
    ws["A1"].font = title_font
    ws["A2"] = "PRESUPUESTO TÉCNICO Y ECONÓMICO: P-2026-0042"
    ws["A2"].font = subtitle_font

    ws["A4"] = "Cliente:"
    ws["B4"] = "Comunidad de Propietarios / Dirección Técnica de Infraestructuras"
    ws["A5"] = "Ubicación:"
    ws["B5"] = "Cámara subterránea / Pozo técnico de registro bajo calzada - Menorca"
    ws["A6"] = "Fecha:"
    ws["B6"] = "23/10/2026"
    ws["A7"] = "Concepto:"
    ws["B7"] = "Sustitución de escalera de servicio en acero galvanizado en caliente y peldaños de tramex (H = 5,73 m)"

    for r in range(4, 8):
        ws[f"A{r}"].font = bold_font
        ws[f"B{r}"].font = regular_font

    # Resumen de Especificaciones Técnicas
    ws["A9"] = "PARÁMETROS TÉCNICOS"
    ws["A9"].font = bold_font
    ws["A9"].fill = section_fill
    ws["B9"] = "Desnivel total: 5,73 m | Desarrollo en planta: 3,48 m | Boca acceso: 161 x 53 cm | 2 tramos + 1 descansillo | 26 peldaños Tramex"
    ws["B9"].font = regular_font
    ws["B9"].fill = section_fill

    # Tabla de partidas
    headers = ["Partida", "Concepto y Descripción Técnica", "Unidad", "Cantidad", "Precio Unitario (€)", "Importe (€)"]
    start_row = 11

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx not in [2] else "left", vertical="center")

    data = [
        ["01", "Desmontaje, oxicorte y retirada de escalera existente degradada (incl. izado por boca superior 161x53 cm, gestión autorizada y seguridad espacio confinado)", "lote", 1, 1450.00],
        ["02", "Fabricación en taller de estructura portante, zancas y descansillo en acero S275JR (4 zancas UPN 140, cartelas soldadas, pilares 80x80x4, barandillas con rodapié s/ UNE-EN ISO 14122-3)", "lote", 1, 3650.00],
        ["03", "Tratamiento anticorrosivo de Galvanizado en Caliente por inmersión según UNE-EN ISO 1461 (70-85 µm) frente a ambiente húmedo C4/C5, portes taller-galvanizadora incluidos", "lote", 1, 1280.00],
        ["04", "Peldaños normalizados y meseta en rejilla electrofundida Tramex galvanizada (26 peldaños con cantonera perforada antideslizante, meseta 140x70 cm y tornillería cincada/Inox A4)", "lote", 1, 2180.00],
        ["05", "Transporte a obra, maniobras con camión grúa pluma en vía pública para descenso vertical por trampilla de 161x53 cm e instalación técnica con anclajes químicos Inox A4", "lote", 1, 2650.00]
    ]

    for i, row in enumerate(data, start=start_row + 1):
        ws.cell(row=i, column=1, value=row[0]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=2, value=row[1]).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=i, column=3, value=row[2]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=4, value=row[3]).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=i, column=5, value=row[4]).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=i, column=6, value=f"=D{i}*E{i}").alignment = Alignment(horizontal="right", vertical="top")

        for col_idx in range(1, 7):
            c = ws.cell(row=i, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx in [4]:
                c.number_format = '#,##0.00'
            elif col_idx in [5, 6]:
                c.number_format = '#,##0.00 €'

    last_data_row = start_row + len(data)

    # Totales
    r_bi = last_data_row + 2
    ws.cell(row=r_bi, column=5, value="Base Imponible:").font = bold_font
    ws.cell(row=r_bi, column=5).alignment = Alignment(horizontal="right")
    cell_bi = ws.cell(row=r_bi, column=6, value=f"=SUM(F{start_row+1}:F{last_data_row})")
    cell_bi.font = bold_font
    cell_bi.number_format = '#,##0.00 €'

    r_iva = r_bi + 1
    ws.cell(row=r_iva, column=5, value="I.V.A. (21%):").font = bold_font
    ws.cell(row=r_iva, column=5).alignment = Alignment(horizontal="right")
    cell_iva = ws.cell(row=r_iva, column=6, value=f"=F{r_bi}*0.21")
    cell_iva.font = bold_font
    cell_iva.number_format = '#,##0.00 €'

    r_tot = r_iva + 1
    ws.cell(row=r_tot, column=5, value="TOTAL PRESUPUESTO:").font = accent_total_font
    ws.cell(row=r_tot, column=5).alignment = Alignment(horizontal="right")
    ws.cell(row=r_tot, column=5).border = double_bottom_border
    cell_tot = ws.cell(row=r_tot, column=6, value=f"=F{r_bi}+F{r_iva}")
    cell_tot.font = accent_total_font
    cell_tot.number_format = '#,##0.00 €'
    cell_tot.border = double_bottom_border

    # Bloque de Condiciones
    r_cond = r_tot + 2
    ws.cell(row=r_cond, column=1, value="CONDICIONES COMERCIALES Y GARANTÍAS").font = bold_font
    ws.cell(row=r_cond+1, column=1, value="• Forma de pago:").font = bold_font
    ws.cell(row=r_cond+1, column=2, value="40% a la confirmación del encargo, 30% a la salida a planta galvanizadora y 30% a la recepción conforme en obra").font = regular_font
    ws.cell(row=r_cond+2, column=1, value="• Garantía:").font = bold_font
    ws.cell(row=r_cond+2, column=2, value="5 años en recubrimiento galvanizado según UNE-EN ISO 1461 y 2 años en estructura y montaje").font = regular_font
    ws.cell(row=r_cond+3, column=1, value="• Plazo de entrega:").font = bold_font
    ws.cell(row=r_cond+3, column=2, value="3 a 4 semanas para taller y galvanizado; 2 a 3 días laborables para montaje en obra").font = regular_font
    ws.cell(row=r_cond+4, column=1, value="• Validez:").font = bold_font
    ws.cell(row=r_cond+4, column=2, value="30 días naturales desde la fecha de emisión").font = regular_font

    # Anchos de columna
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 68
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 22
    ws.column_dimensions['F'].width = 20

    xlsx_path = r"c:\Users\pablo\Documents\Presupuestador\presupuestos\P-2026-0042-escalera-galvanizada-peldanos-tramex\presupuesto.xlsx"
    wb.save(xlsx_path)
    print(f"Excel guardado exitosamente en: {xlsx_path}")

def generate_pdf():
    html_path = r"c:\Users\pablo\Documents\Presupuestador\presupuestos\P-2026-0042-escalera-galvanizada-peldanos-tramex\presupuesto-cliente.html"
    pdf_path = r"c:\Users\pablo\Documents\Presupuestador\presupuestos\P-2026-0042-escalera-galvanizada-peldanos-tramex\presupuesto-cliente.pdf"
    
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(f"file:///{html_path.replace(os.sep, '/')}")
        page.pdf(
            path=pdf_path,
            format="A4",
            margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
            print_background=True
        )
        browser.close()
    print(f"PDF generado exitosamente en: {pdf_path}")

if __name__ == "__main__":
    create_excel()
    generate_pdf()
