import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.sync_api import sync_playwright

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "P-2026-0045 Escalera L Iroco"

    # Paleta de colores cálida/técnica (Grafito / Marrón Ocre Dorado / Slate)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=11, bold=True, color="B45309")
    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    accent_total_font = Font(name="Segoe UI", size=12, bold=True, color="B45309")
    section_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    double_bottom_border = Border(
        top=Side(style='thin', color='0F172A'),
        bottom=Side(style='double', color='0F172A')
    )

    # Cabecera Presupuesto
    ws["A1"] = "HAM ESTRUCTURAS METÁLICAS"
    ws["A1"].font = title_font
    ws["A2"] = "PRESUPUESTO TÉCNICO Y ECONÓMICO: P-2026-0045 (REV. 3)"
    ws["A2"].font = subtitle_font

    ws["A4"] = "Cliente:"
    ws["B4"] = "Costel Pindea (Constructora)"
    ws["A5"] = "Ubicación:"
    ws["B5"] = "Obra en Menorca - Islas Baleares"
    ws["A6"] = "Fecha:"
    ws["B6"] = "28/10/2026"
    ws["A7"] = "Concepto:"
    ws["B7"] = "Escalera en L: base tubular arranque (3 escalones), tramo volado pletina mural, forrado Iroko y varillas a viga de madera (fijaciones a definir)"

    for r in range(4, 8):
        ws[f"A{r}"].font = bold_font
        ws[f"B{r}"].font = regular_font

    # Resumen de Especificaciones Técnicas
    ws["A9"] = "PARÁMETROS TÉCNICOS"
    ws["A9"].font = bold_font
    ws["A9"].fill = section_fill
    ws["B9"] = "Escalera en L | Base tubular para 3 escalones de arranque | Zanca pletina 150x12 mm amurada a pared | Tubos 50x50x5 mm | Varillas Ø10 mm a viga de madera | Forrado 19 peldaños en Iroko macizo | Detalles de tuercas y fijaciones a definir"
    ws["B9"].font = regular_font
    ws["B9"].fill = section_fill

    # Tabla de partidas
    headers = ["Partida", "Concepto y Descripción Técnica", "Unidad", "Cantidad", "Precio Unitario (€)", "Importe (€)"]
    start_row = 11

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx not in [2] else "left", vertical="center")

    data = [
        ["01", "Toma de medidas, replanteo láser in situ y desarrollo de planos de despiece para configuración en L (giro a 90° en arranque, verificación de plomos, cota de viga y definición de detalles de tuercas y fijaciones)", "lote", 1, 572.00],
        ["02", "Fabricación de base tubular estructural para arranque en L de 3 escalones (chasis soldado en tubo 80x40x3 y 50x50x3 mm S275JR, pletinas de anclaje a solera y encuentro con zanca mural; fijaciones a definir)", "lote", 1, 1265.00],
        ["03", "Fabricación en taller de zanca pletina mural (150x12 mm S275JR) y 16 ménsulas en tubo cuadrado estructural de 50x50x5 mm soldadas, con tuercas remache insertadas (tipo y métrica a definir)", "lote", 1, 4235.00],
        ["04", "Barandilla vertical suspendida de varillas calibradas/roscadas Ø10 mm (32 uds para hueco < 10 cm) y perfil angular superior L 60x40x5 mm fijado a viga de madera (detalles de tuercas y fijaciones a definir)", "lote", 1, 2145.00],
        ["05", "Tratamiento anticorrosivo y lacado de toda la cerrajería en taller: limpieza, imprimación epoxi fosfatante antioxidante y esmalte poliuretano secado al horno (RAL 7016 / 9005 texturado)", "lote", 1, 1078.00],
        ["06", "Forrado integral de 19 peldaños en Madera Maciza de Iroko de 1ª calidad seleccionada (3 en base arranque + 16 en tramo volado; huella 38 mm, frentes 20 mm y 3 capas barniz poliuretano tránsito)", "lote", 1, 5478.00],
        ["07", "Transporte a obra en Menorca, fijaciones a solera, muro y viga de madera, montaje de varillas, pegado elastomérico de peldaños de Iroko y remates (sistemas de fijación y tuercas a definir)", "lote", 1, 3058.00]
    ]

    for i, row in enumerate(data, start=start_row + 1):
        ws.cell(row=i, column=1, value=row[0]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=2, value=row[1]).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=i, column=3, value=row[2]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=4, value=row[3]).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=i, column=5, value=row[4]).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=i, column=6, value=f"=D{i}*E{i}").alignment = Alignment(horizontal="right", vertical="top")

        for col_idx in range(1, 7):
            c = ws.cell(row=i, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx in [4]:
                c.number_format = '#,##0.00'
            elif col_idx in [5, 6]:
                c.number_format = '#,##0.00 €'

    last_data_row = start_row + len(data)

    # Totales
    r_bi = last_data_row + 2
    ws.cell(row=r_bi, column=5, value="Base Imponible:").font = bold_font
    ws.cell(row=r_bi, column=5).alignment = Alignment(horizontal="right")
    cell_bi = ws.cell(row=r_bi, column=6, value=f"=SUM(F{start_row+1}:F{last_data_row})")
    cell_bi.font = bold_font
    cell_bi.number_format = '#,##0.00 €'

    r_iva = r_bi + 1
    ws.cell(row=r_iva, column=5, value="I.V.A. (21%):").font = bold_font
    ws.cell(row=r_iva, column=5).alignment = Alignment(horizontal="right")
    cell_iva = ws.cell(row=r_iva, column=6, value=f"=F{r_bi}*0.21")
    cell_iva.font = bold_font
    cell_iva.number_format = '#,##0.00 €'

    r_tot = r_iva + 1
    ws.cell(row=r_tot, column=5, value="TOTAL PRESUPUESTO:").font = accent_total_font
    ws.cell(row=r_tot, column=5).alignment = Alignment(horizontal="right")
    ws.cell(row=r_tot, column=5).border = double_bottom_border
    cell_tot = ws.cell(row=r_tot, column=6, value=f"=F{r_bi}+F{r_iva}")
    cell_tot.font = accent_total_font
    cell_tot.number_format = '#,##0.00 €'
    cell_tot.border = double_bottom_border

    # Bloque de Condiciones
    r_cond = r_tot + 2
    ws.cell(row=r_cond, column=1, value="CONDICIONES COMERCIALES Y GARANTÍAS").font = bold_font
    ws.cell(row=r_cond+1, column=1, value="• Forma de pago:").font = bold_font
    ws.cell(row=r_cond+1, column=2, value="100% a la aceptación del presupuesto").font = bold_font
    ws.cell(row=r_cond+2, column=1, value="• Fijaciones / tuercas:").font = bold_font
    ws.cell(row=r_cond+2, column=2, value="Detalles constructivos de tuercas y sistemas de fijación a muro, solera y viga a definir en replanteo").font = regular_font
    ws.cell(row=r_cond+3, column=1, value="• Garantía:").font = bold_font
    ws.cell(row=r_cond+3, column=2, value="2 años en estructura de acero, cerrajería, fijaciones y acabados según normativa vigente").font = regular_font
    ws.cell(row=r_cond+4, column=1, value="• Plazo de entrega:").font = bold_font
    ws.cell(row=r_cond+4, column=2, value="4 a 5 semanas para taller de cerrajería y ebanistería; montaje en obra de 4 jornadas laborables").font = regular_font
    ws.cell(row=r_cond+5, column=1, value="• Validez:").font = bold_font
    ws.cell(row=r_cond+5, column=2, value="30 días naturales desde la fecha de emisión").font = regular_font

    # Anchos de columna
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 75
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 22
    ws.column_dimensions['F'].width = 20

    base_dir = os.path.dirname(os.path.abspath(__file__))
    xlsx_path = os.path.join(base_dir, "presupuesto.xlsx")
    wb.save(xlsx_path)
    print(f"Excel guardado exitosamente en: {xlsx_path}")

def generate_pdf():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    html_path = os.path.join(base_dir, "presupuesto-cliente.html")
    pdf_path = os.path.join(base_dir, "presupuesto-cliente.pdf")
    
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(f"file:///{html_path.replace(os.sep, '/')}")
        page.pdf(
            path=pdf_path,
            format="A4",
            margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
            print_background=True
        )
        browser.close()
    print(f"PDF generado exitosamente en: {pdf_path}")

if __name__ == "__main__":
    create_excel()
    generate_pdf()
