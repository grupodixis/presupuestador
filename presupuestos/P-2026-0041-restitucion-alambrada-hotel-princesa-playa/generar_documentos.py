import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.sync_api import sync_playwright

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "P-2026-0041 Princesa Playa"

    # Estilos
    header_fill = PatternFill(start_color="1B3B2B", end_color="1B3B2B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1B3B2B")
    subtitle_font = Font(name="Calibri", size=11, bold=True, color="4B5563")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    accent_total_font = Font(name="Calibri", size=12, bold=True, color="166534")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    double_bottom_border = Border(
        top=Side(style='thin', color='1B3B2B'),
        bottom=Side(style='double', color='1B3B2B')
    )

    # Cabecera Presupuesto
    ws["A1"] = "HAM ESTRUCTURAS METÁLICAS"
    ws["A1"].font = title_font
    ws["A2"] = "PRESUPUESTO TÉCNICO Y ECONÓMICO: P-2026-0041"
    ws["A2"].font = subtitle_font

    ws["A4"] = "Cliente:"
    ws["B4"] = "Hotel Princesa Playa"
    ws["A5"] = "Ubicación:"
    ws["B5"] = "Son Xoriguer / Cala'n Bosch, Ciutadella de Menorca"
    ws["A6"] = "Fecha:"
    ws["B6"] = "23/10/2026"
    ws["A7"] = "Concepto:"
    ws["B7"] = "Restitución de 80 ML de alambrada simple torsión (h=3,00 m, verde RAL 6005)"

    for r in range(4, 8):
        ws[f"A{r}"].font = bold_font
        ws[f"B{r}"].font = regular_font

    # Tabla de partidas
    headers = ["Partida", "Concepto y Descripción", "Unidad", "Cantidad", "Precio Unitario (€)", "Importe (€)"]
    start_row = 9

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx not in [2] else "left", vertical="center")

    data = [
        ["01", "Desmontaje, retirada y gestión de cerramiento previo deteriorado en vertedero autorizado", "ml", 80, 35.00],
        ["02", "Postes tubulares metálicos galvanizados y plastificados verde h=3,00 m (intermedios, esquinas y tornapuntas)", "ml", 80, 75.00],
        ["03", "Malla simple torsión 50/14 plastificada verde h=3,00 m, 4 líneas tensoras, tensores M8 y accesorios", "ml", 80, 65.00],
        ["04", "Mano de obra especializada en montaje, alineación, tesado y remates de seguridad en obra", "ml", 80, 35.00]
    ]

    for i, row in enumerate(data, start=start_row + 1):
        ws.cell(row=i, column=1, value=row[0]).alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=2, value=row[1]).alignment = Alignment(horizontal="left", wrap_text=True)
        ws.cell(row=i, column=3, value=row[2]).alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=4, value=row[3]).alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=5, value=row[4]).alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=6, value=f"=D{i}*E{i}").alignment = Alignment(horizontal="right")

        for col_idx in range(1, 7):
            c = ws.cell(row=i, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx in [4]:
                c.number_format = '#,##0.00'
            elif col_idx in [5, 6]:
                c.number_format = '#,##0.00 €'

    last_data_row = start_row + len(data)

    # Totales
    r_bi = last_data_row + 2
    ws.cell(row=r_bi, column=5, value="Base Imponible:").font = bold_font
    ws.cell(row=r_bi, column=5).alignment = Alignment(horizontal="right")
    cell_bi = ws.cell(row=r_bi, column=6, value=f"=SUM(F{start_row+1}:F{last_data_row})")
    cell_bi.font = bold_font
    cell_bi.number_format = '#,##0.00 €'

    r_iva = r_bi + 1
    ws.cell(row=r_iva, column=5, value="I.V.A. (21%):").font = bold_font
    ws.cell(row=r_iva, column=5).alignment = Alignment(horizontal="right")
    cell_iva = ws.cell(row=r_iva, column=6, value=f"=F{r_bi}*0.21")
    cell_iva.font = bold_font
    cell_iva.number_format = '#,##0.00 €'

    r_tot = r_iva + 1
    ws.cell(row=r_tot, column=5, value="TOTAL PRESUPUESTO:").font = accent_total_font
    ws.cell(row=r_tot, column=5).alignment = Alignment(horizontal="right")
    ws.cell(row=r_tot, column=5).border = double_bottom_border
    cell_tot = ws.cell(row=r_tot, column=6, value=f"=F{r_bi}+F{r_iva}")
    cell_tot.font = accent_total_font
    cell_tot.number_format = '#,##0.00 €'
    cell_tot.border = double_bottom_border

    # Bloque de Condiciones
    r_cond = r_tot + 2
    ws.cell(row=r_cond, column=1, value="CONDICIONES COMERCIALES").font = bold_font
    ws.cell(row=r_cond+1, column=1, value="• Forma de pago:").font = bold_font
    ws.cell(row=r_cond+1, column=2, value="80% a la aceptación del presupuesto (16.262,40 € IVA incl.) y 20% a la finalización (4.065,60 € IVA incl.)").font = regular_font
    ws.cell(row=r_cond+2, column=1, value="• Garantía:").font = bold_font
    ws.cell(row=r_cond+2, column=2, value="1 año en materiales y montaje frente a defectos de fabricación o instalación").font = regular_font
    ws.cell(row=r_cond+3, column=1, value="• Plazo de ejecución:").font = bold_font
    ws.cell(row=r_cond+3, column=2, value="4 a 6 días laborables tras replanteo y recepción de acopios").font = regular_font
    ws.cell(row=r_cond+4, column=1, value="• Validez:").font = bold_font
    ws.cell(row=r_cond+4, column=2, value="30 días naturales desde la fecha de emisión").font = regular_font

    # Anchos de columna
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 65
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 22
    ws.column_dimensions['F'].width = 20

    xlsx_path = r"c:\Users\pablo\Documents\Presupuestador\presupuestos\P-2026-0041-restitucion-alambrada-hotel-princesa-playa\presupuesto.xlsx"
    wb.save(xlsx_path)
    print(f"Excel guardado exitosamente en: {xlsx_path}")

def generate_pdf():
    html_path = r"c:\Users\pablo\Documents\Presupuestador\presupuestos\P-2026-0041-restitucion-alambrada-hotel-princesa-playa\presupuesto-cliente.html"
    pdf_path = r"c:\Users\pablo\Documents\Presupuestador\presupuestos\P-2026-0041-restitucion-alambrada-hotel-princesa-playa\presupuesto-cliente.pdf"
    
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(f"file:///{html_path.replace(os.sep, '/')}")
        page.pdf(
            path=pdf_path,
            format="A4",
            margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
            print_background=True
        )
        browser.close()
    print(f"PDF generado exitosamente en: {pdf_path}")

if __name__ == "__main__":
    create_excel()
    generate_pdf()
