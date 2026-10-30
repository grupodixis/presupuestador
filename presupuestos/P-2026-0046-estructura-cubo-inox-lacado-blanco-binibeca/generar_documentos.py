import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.sync_api import sync_playwright

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "P-2026-0046 Cubo Inox Blanco"

    # Estilos y Paleta de Colores Corporativa
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=11, bold=True, color="0284C7")
    regular_font = Font(name="Segoe UI", size=9.5)
    bold_font = Font(name="Segoe UI", size=9.5, bold=True)
    accent_total_font = Font(name="Segoe UI", size=11, bold=True, color="0284C7")
    section_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Cabecera Presupuesto
    ws["A1"] = "HAM ESTRUCTURAS METÁLICAS MENORCA"
    ws["A1"].font = title_font
    ws["A2"] = "PRESUPUESTO TÉCNICO Y VALORACIÓN ECONÓMICA: P-2026-0046 (REV. 3)"
    ws["A2"].font = subtitle_font

    ws["A4"] = "Cliente:"
    ws["B4"] = "Teresa Echezarreta Codina"
    ws["A5"] = "Ubicación:"
    ws["B5"] = "c/ Llaut 59, Urb. Binibeca Vell, Sant Lluís, Menorca"
    ws["A6"] = "Proyecto / Arq.:"
    ws["B6"] = "Cristina Moll Taltavull · Expte: 2024 13_MOD (Reforma Unifamiliar con Piscina)"
    ws["A7"] = "Fecha y Validez:"
    ws["B7"] = "30/10/2026 · Validez 30 días"
    ws["A8"] = "Concepto:"
    ws["B8"] = "Estructura cúbica minimalista en tubo de Acero Inoxidable 40x40x2 mm lacado blanco, fijada sobre zapatas ejecutadas por constructor"

    for r in range(4, 9):
        ws[f"A{r}"].font = bold_font
        ws[f"B{r}"].font = regular_font

    # Resumen Parámetros Técnicos
    ws["A10"] = "PARÁMETROS TÉCNICOS"
    ws["A10"].font = bold_font
    ws["A10"].fill = section_fill
    ws["B10"] = "Medidas: 4,00 x 3,00 x 2,50 m (12 m²) | Geometría cúbica ortoédrica | Tubo cuadrado de Acero Inoxidable 40x40x2 mm | Soldadura TIG pulida | Termolacado al horno en Blanco exterior (RAL 9010/9016) | Obra civil y zapatas a cargo del constructor | Transporte y montaje de estructura fijada sobre puntos de obra"
    ws["B10"].font = regular_font
    ws["B10"].fill = section_fill

    # Tabla de partidas
    headers = ["Partida", "Concepto y Descripción Técnica", "Unidad", "Cantidad", "Precio Unitario (€)", "Importe (€)"]
    start_row = 12

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx not in [2] else "left", vertical="center")

    data = [
        ["01", "Replanteo técnico in situ, nivelación y desarrollo de planos de taller: visita a la obra en Binibeca Vell para comprobación de cotas y coordinación de apoyos con el constructor previa ejecución de zapatas/solera, y elaboración de planos de taller", "lote", 1, 320.00],
        ["02", "Fabricación en taller de estructura cúbica en tubo de Acero Inoxidable 40x40x2 mm: 4 pilares verticales de 2,50 m, marco perimetral superior (2 vigas de 4 m + 2 vigas de 3 m) y costillas intermedias superiores transversales de rigidización. Cortes ingleteados a 45º, soldadura continua TIG en atmósfera inerte de argón y aristas esmeriladas para acabado monolítico", "lote", 1, 2150.00],
        ["03", "Tratamiento superficial de termolacado al horno en Blanco Exterior: desengrasado, acondicionamiento superficial para acero inoxidable, imprimación de adherencia técnica y termolacado al horno con pintura en polvo de resinas de poliéster arquitectónico (Blanco RAL 9010/9016) resistente a radiación UV y salinidad marina", "lote", 1, 680.00],
        ["04", "Transporte a Binibeca Vell y montaje de estructura sobre puntos preparados: transporte especializado a obra, descarga, elevación de módulos, aplomado y fijación estructural sobre los puntos/zapatas ejecutados por el constructor, verificación de niveles, repaso de lacado y retirada de embalajes", "lote", 1, 980.00]
    ]

    for i, row in enumerate(data, start=start_row + 1):
        ws.cell(row=i, column=1, value=row[0]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=2, value=row[1]).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=i, column=3, value=row[2]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=4, value=row[3]).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=i, column=5, value=row[4]).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=i, column=6, value=f"=D{i}*E{i}").alignment = Alignment(horizontal="right", vertical="top")

        for col_idx in range(1, 7):
            c = ws.cell(row=i, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx in [4]:
                c.number_format = '#,##0.00'
            elif col_idx in [5, 6]:
                c.number_format = '#,##0.00 €'

    last_data_row = start_row + len(data)

    # Totales
    r_bi = last_data_row + 2
    ws.cell(row=r_bi, column=5, value="Base Imponible:").font = bold_font
    ws.cell(row=r_bi, column=6, value=f"=SUM(F{start_row+1}:F{last_data_row})").font = bold_font
    ws.cell(row=r_bi, column=6).number_format = '#,##0.00 €'

    r_iva = r_bi + 1
    ws.cell(row=r_iva, column=5, value="IVA (21%):").font = regular_font
    ws.cell(row=r_iva, column=6, value=f"=F{r_bi}*0.21").font = regular_font
    ws.cell(row=r_iva, column=6).number_format = '#,##0.00 €'

    r_tot = r_iva + 1
    ws.cell(row=r_tot, column=5, value="TOTAL PRESUPUESTO:").font = accent_total_font
    ws.cell(row=r_tot, column=6, value=f"=F{r_bi}+F{r_iva}").font = accent_total_font
    ws.cell(row=r_tot, column=6).number_format = '#,##0.00 €'

    # Variante 4x2 m
    r_var = r_tot + 2
    ws.cell(row=r_var, column=2, value="VARIANTE ALTERNATIVA: Estructura Cúbica 4,00 x 2,00 m (8,00 m²)").font = bold_font
    ws.cell(row=r_var, column=5, value="Base Imponible:").font = regular_font
    ws.cell(row=r_var, column=6, value=3500.00).font = bold_font
    ws.cell(row=r_var, column=6).number_format = '#,##0.00 €'
    ws.cell(row=r_var+1, column=5, value="Total con IVA (21%):").font = regular_font
    ws.cell(row=r_var+1, column=6, value=4235.00).font = bold_font
    ws.cell(row=r_var+1, column=6).number_format = '#,##0.00 €'

    # Ajuste de anchos de columna
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 62
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 18

    excel_path = os.path.join(os.path.dirname(__file__), "presupuesto.xlsx")
    wb.save(excel_path)
    print(f"Excel generado con exito: {excel_path}")
    return excel_path

def convert_html_to_pdf():
    base_dir = os.path.dirname(__file__)
    cliente_html = os.path.abspath(os.path.join(base_dir, "presupuesto-cliente.html"))
    cliente_pdf = os.path.abspath(os.path.join(base_dir, "presupuesto-cliente.pdf"))
    tecnico_html = os.path.abspath(os.path.join(base_dir, "presupuesto-interno-tecnico.html"))
    tecnico_pdf = os.path.abspath(os.path.join(base_dir, "presupuesto-interno-tecnico.pdf"))

    with sync_playwright() as p:
        browser = p.chromium.launch()
        
        # Cliente
        page = browser.new_page()
        page.goto(f"file:///{cliente_html.replace(os.sep, '/')}", wait_until="networkidle")
        page.pdf(
            path=cliente_pdf,
            format="A4",
            print_background=True,
            margin={"top": "0mm", "bottom": "0mm", "left": "0mm", "right": "0mm"}
        )
        print(f"PDF Cliente generado: {cliente_pdf}")
        
        # Tecnico
        page2 = browser.new_page()
        page2.goto(f"file:///{tecnico_html.replace(os.sep, '/')}", wait_until="networkidle")
        page2.pdf(
            path=tecnico_pdf,
            format="A4",
            print_background=True,
            margin={"top": "0mm", "bottom": "0mm", "left": "0mm", "right": "0mm"}
        )
        print(f"PDF Tecnico generado: {tecnico_pdf}")
        
        browser.close()

if __name__ == "__main__":
    create_excel()
    convert_html_to_pdf()
