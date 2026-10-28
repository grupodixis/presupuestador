import os
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.sync_api import sync_playwright

def load_data():
    json_path = os.path.join(os.path.dirname(__file__), "datos.json")
    with open(json_path, "r", encoding="utf-8") as f:
        return json.load(f)

def create_excel():
    data = load_data()
    wb = Workbook()

    # -------------------------------------------------------------
    # HOJA 1: PRESUPUESTO CLIENTE
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Presupuesto Cliente"

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_optional_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    regular_font = Font(name="Segoe UI", size=9.5, color="1E293B")
    accent_total_font = Font(name="Segoe UI", size=10.5, bold=True, color="B91C1C")
    subtotal_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    optional_subtotal_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    grand_total_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    section_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    thin_border_side = Side(style='thin', color="CBD5E1")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    double_bottom_border = Border(top=thin_border_side, bottom=Side(style='double', color="0F172A"))

    ws1["A1"] = "HAM ESTRUCTURAS METÁLICAS (DIXIS RECURSOS Y SOLUCIONES, SL)"
    ws1["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    ws1["A2"] = f"PRESUPUESTO TÉCNICO OFICIAL: {data['numeroPresupuesto']} | REF: {data['referencia']}"
    ws1["A2"].font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")

    ws1["A4"] = "Cliente:"
    ws1["B4"] = f"{data['cliente']['empresa']} | Atn: {data['cliente']['atencion']}"
    ws1["A5"] = "Proyecto:"
    ws1["B5"] = f"{data['cliente']['proyecto']} | Promotor: {data['cliente']['promotorFinal']}"
    ws1["A6"] = "Emplazamiento:"
    ws1["B6"] = f"{data['cliente']['emplazamiento']}, {data['cliente']['poblacion']} ({data['cliente']['isla']})"
    ws1["A7"] = "Fecha de emisión:"
    ws1["B7"] = "28/10/2026"

    for r in range(4, 8):
        ws1[f"A{r}"].font = bold_font
        ws1[f"B{r}"].font = regular_font

    ws1["A9"] = "SISTEMA TÉCNICO"
    ws1["A9"].font = bold_font
    ws1["A9"].fill = section_fill
    ws1["B9"] = "Panel Sándwich Cubierta PIR 40 mm (B-s1,d0) + Chapas grecadas voladizos y tabica con acabado Granite® HDX55 (55 µm) por AMBAS CARAS | Fijaciones roscachapa galvanizadas con arandela vulcanizada EPDM | Medios mecánicos y fletes integrados | Opcionales: 2 escalas de gato inox con crinolina (+30%) + Línea LuxTop® (18.000 €)"
    ws1["B9"].font = regular_font
    ws1["B9"].fill = section_fill

    # Tabla de Sectores
    ws1["A11"] = "DESGLOSE DE SUPERFICIES DE ACTUACIÓN"
    ws1["A11"].font = bold_font
    sector_headers = ["Zona", "Descripción del Elemento", "Tipología Constructiva", "Superficie (m²)", "Acabado de Protección", "Ambiente"]
    for c_idx, s_head in enumerate(sector_headers, 1):
        cell = ws1.cell(row=12, column=c_idx, value=s_head)
        cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center" if c_idx not in [2, 3] else "left", vertical="center")

    sectores_data = [
        ["1", "Cubierta faldones principales a dos aguas", "Panel sándwich 40 mm PIR (UNE-EN 14509)", 1042.50, "Granite® HDX55 (2 caras)", "Marino C5-M / Aéreo"],
        ["2", "Cubierta faldón a un agua", "Panel sándwich 40 mm PIR adaptado a petos", 328.50, "Granite® HDX55 (2 caras)", "Marino C5-M / Aéreo"],
        ["3", "Voladizos en fachadas (cobertura superior)", "Chapa/panel grecado sobre ménsulas metálicas", 533.00, "Granite® HDX55 (2 caras) + Imprimación", "Marino C5-M / Aéreo"],
        ["4", "Lado inferior de voladizos (cielo raso)", "Chapa grecada bajo viguetas en ménsula", 533.60, "Granite® HDX55 (2 caras)", "Marino C5-M / Aéreo"],
        ["5", "Tabica vertical de fachada sobre marquesinas", "Chapa grecada en paramento vertical", 186.50, "Granite® HDX55 (2 caras)", "Marino C5-M / Aéreo"]
    ]

    for s_idx, s_row in enumerate(sectores_data, start=13):
        for col_i in range(1, 7):
            c = ws1.cell(row=s_idx, column=col_i, value=s_row[col_i-1])
            c.font = regular_font
            c.border = thin_border
            if col_i == 4:
                c.number_format = '#,##0.00'
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col_i in [1, 5, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center")

    row_tot_sect = 13 + len(sectores_data)
    ws1.cell(row=row_tot_sect, column=1, value="")
    ws1.cell(row=row_tot_sect, column=2, value="TOTAL SUPERFICIE DE ACTUACIÓN:").font = bold_font
    ws1.cell(row=row_tot_sect, column=3, value="")
    c_tot_m2 = ws1.cell(row=row_tot_sect, column=4, value=f"=SUM(D13:D{row_tot_sect-1})")
    c_tot_m2.font = bold_font
    c_tot_m2.number_format = '#,##0.00'
    c_tot_m2.alignment = Alignment(horizontal="right", vertical="center")
    ws1.cell(row=row_tot_sect, column=5, value="m² totales")
    ws1.cell(row=row_tot_sect, column=6, value="")
    for col_i in range(1, 7):
        ws1.cell(row=row_tot_sect, column=col_i).border = thin_border
        ws1.cell(row=row_tot_sect, column=col_i).fill = section_fill

    # Tabla de Partidas Base
    start_row = row_tot_sect + 2
    ws1.cell(row=start_row, column=1, value="VALORACIÓN ECONÓMICA DE PARTIDAS BASE (CON MEDIOS INTEGRADOS)").font = bold_font
    headers = ["Partida", "Concepto y Descripción Técnica", "Unidad", "Cantidad", "Precio Unitario (€)", "Importe (€)"]
    start_row += 1

    for col_idx, header in enumerate(headers, 1):
        cell = ws1.cell(row=start_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx not in [2] else "left", vertical="center")

    partidas = data["partidas"]
    for i, p in enumerate(partidas, start=start_row + 1):
        ws1.cell(row=i, column=1, value=p["id"]).alignment = Alignment(horizontal="center", vertical="top")
        desc_text = f"{p['concepto']}\n{p['descripcion']}"
        ws1.cell(row=i, column=2, value=desc_text).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws1.cell(row=i, column=3, value=p["unidad"]).alignment = Alignment(horizontal="center", vertical="top")
        ws1.cell(row=i, column=4, value=p["cantidad"]).alignment = Alignment(horizontal="right", vertical="top")
        ws1.cell(row=i, column=5, value=p["precioUnitario"]).alignment = Alignment(horizontal="right", vertical="top")
        ws1.cell(row=i, column=6, value=f"=D{i}*E{i}").alignment = Alignment(horizontal="right", vertical="top")

        for col_idx in range(1, 7):
            c = ws1.cell(row=i, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx in [4]:
                c.number_format = '#,##0.00'
            elif col_idx in [5, 6]:
                c.number_format = '#,##0.00 €'

    last_base_row = start_row + len(partidas)

    # Subtotales Base
    r_bi_base = last_base_row + 1
    ws1.cell(row=r_bi_base, column=5, value="Subtotal Base Imponible (Obra Base):").font = bold_font
    ws1.cell(row=r_bi_base, column=5).alignment = Alignment(horizontal="right")
    cell_bi_base = ws1.cell(row=r_bi_base, column=6, value=f"=SUM(F{start_row+1}:F{last_base_row})")
    cell_bi_base.font = bold_font
    cell_bi_base.number_format = '#,##0.00 €'
    ws1.cell(row=r_bi_base, column=5).fill = subtotal_fill
    cell_bi_base.fill = subtotal_fill

    r_iva_base = r_bi_base + 1
    ws1.cell(row=r_iva_base, column=5, value="I.V.A. (21%) Obra Base:").font = regular_font
    ws1.cell(row=r_iva_base, column=5).alignment = Alignment(horizontal="right")
    cell_iva_base = ws1.cell(row=r_iva_base, column=6, value=f"=F{r_bi_base}*0.21")
    cell_iva_base.font = regular_font
    cell_iva_base.number_format = '#,##0.00 €'

    r_tot_base = r_iva_base + 1
    ws1.cell(row=r_tot_base, column=5, value="TOTAL OBRA BASE (IVA incl.):").font = accent_total_font
    ws1.cell(row=r_tot_base, column=5).alignment = Alignment(horizontal="right")
    cell_tot_base = ws1.cell(row=r_tot_base, column=6, value=f"=F{r_bi_base}+F{r_iva_base}")
    cell_tot_base.font = accent_total_font
    cell_tot_base.number_format = '#,##0.00 €'
    ws1.cell(row=r_tot_base, column=5).fill = subtotal_fill
    cell_tot_base.fill = subtotal_fill
    ws1.cell(row=r_tot_base, column=5).border = double_bottom_border
    cell_tot_base.border = double_bottom_border

    # BLOQUE DE PARTIDAS OPCIONALES
    r_opc_title = r_tot_base + 2
    ws1.cell(row=r_opc_title, column=1, value="PARTIDAS OPCIONALES INDEPENDIENTES (ACCESOS EN ACERO INOXIDABLE Y LÍNEA LUXTOP)").font = bold_font

    start_opc_row = r_opc_title + 1
    for col_idx, header in enumerate(headers, 1):
        cell = ws1.cell(row=start_opc_row, column=col_idx, value=header)
        cell.fill = header_optional_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx not in [2] else "left", vertical="center")

    opcionales = data["partidasOpcionales"]
    for i, op in enumerate(opcionales, start=start_opc_row + 1):
        ws1.cell(row=i, column=1, value=op["id"]).alignment = Alignment(horizontal="center", vertical="top")
        desc_text = f"{op['concepto']}\n{op['descripcion']}"
        ws1.cell(row=i, column=2, value=desc_text).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws1.cell(row=i, column=3, value=op["unidad"]).alignment = Alignment(horizontal="center", vertical="top")
        ws1.cell(row=i, column=4, value=op["cantidad"]).alignment = Alignment(horizontal="right", vertical="top")
        ws1.cell(row=i, column=5, value=op["precioUnitario"]).alignment = Alignment(horizontal="right", vertical="top")
        ws1.cell(row=i, column=6, value=f"=D{i}*E{i}").alignment = Alignment(horizontal="right", vertical="top")

        for col_idx in range(1, 7):
            c = ws1.cell(row=i, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx in [4]:
                c.number_format = '#,##0.00'
            elif col_idx in [5, 6]:
                c.number_format = '#,##0.00 €'

    last_opc_row = start_opc_row + len(opcionales)

    # Subtotales Opcionales
    r_bi_opc = last_opc_row + 1
    ws1.cell(row=r_bi_opc, column=5, value="Subtotal Opcionales (Base Imp.):").font = bold_font
    ws1.cell(row=r_bi_opc, column=5).alignment = Alignment(horizontal="right")
    cell_bi_opc = ws1.cell(row=r_bi_opc, column=6, value=f"=SUM(F{start_opc_row+1}:F{last_opc_row})")
    cell_bi_opc.font = bold_font
    cell_bi_opc.number_format = '#,##0.00 €'
    ws1.cell(row=r_bi_opc, column=5).fill = optional_subtotal_fill
    cell_bi_opc.fill = optional_subtotal_fill

    r_iva_opc = r_bi_opc + 1
    ws1.cell(row=r_iva_opc, column=5, value="I.V.A. (21%) Opcionales:").font = regular_font
    ws1.cell(row=r_iva_opc, column=5).alignment = Alignment(horizontal="right")
    cell_iva_opc = ws1.cell(row=r_iva_opc, column=6, value=f"=F{r_bi_opc}*0.21")
    cell_iva_opc.font = regular_font
    cell_iva_opc.number_format = '#,##0.00 €'

    r_tot_opc = r_iva_opc + 1
    ws1.cell(row=r_tot_opc, column=5, value="TOTAL OPCIONALES (IVA incl.):").font = bold_font
    ws1.cell(row=r_tot_opc, column=5).alignment = Alignment(horizontal="right")
    cell_tot_opc = ws1.cell(row=r_tot_opc, column=6, value=f"=F{r_bi_opc}+F{r_iva_opc}")
    cell_tot_opc.font = bold_font
    cell_tot_opc.number_format = '#,##0.00 €'
    ws1.cell(row=r_tot_opc, column=5).fill = optional_subtotal_fill
    cell_tot_opc.fill = optional_subtotal_fill
    ws1.cell(row=r_tot_opc, column=5).border = double_bottom_border
    cell_tot_opc.border = double_bottom_border

    # TOTAL GENERAL CONTRATO CON OPCIONALES
    r_grand = r_tot_opc + 2
    ws1.cell(row=r_grand, column=5, value="TOTAL GENERAL CONTRATO (BASE + OPCIONALES):").font = Font(name="Segoe UI", size=11, bold=True, color="065F46")
    ws1.cell(row=r_grand, column=5).alignment = Alignment(horizontal="right")
    cell_grand_bi = ws1.cell(row=r_grand, column=6, value=f"=F{r_bi_base}+F{r_bi_opc}")
    cell_grand_bi.font = Font(name="Segoe UI", size=11, bold=True, color="065F46")
    cell_grand_bi.number_format = '#,##0.00 €'
    ws1.cell(row=r_grand, column=5).fill = grand_total_fill
    cell_grand_bi.fill = grand_total_fill

    r_grand_iva = r_grand + 1
    ws1.cell(row=r_grand_iva, column=5, value="TOTAL I.V.A. CONJUNTO (21%):").font = bold_font
    ws1.cell(row=r_grand_iva, column=5).alignment = Alignment(horizontal="right")
    cell_grand_iva = ws1.cell(row=r_grand_iva, column=6, value=f"=F{r_grand}*0.21")
    cell_grand_iva.font = bold_font
    cell_grand_iva.number_format = '#,##0.00 €'

    r_grand_tot = r_grand_iva + 1
    ws1.cell(row=r_grand_tot, column=5, value="TOTAL GENERAL FACTURA (IVA incl.):").font = Font(name="Segoe UI", size=12, bold=True, color="B91C1C")
    ws1.cell(row=r_grand_tot, column=5).alignment = Alignment(horizontal="right")
    cell_grand_tot = ws1.cell(row=r_grand_tot, column=6, value=f"=F{r_grand}+F{r_grand_iva}")
    cell_grand_tot.font = Font(name="Segoe UI", size=12, bold=True, color="B91C1C")
    cell_grand_tot.number_format = '#,##0.00 €'
    ws1.cell(row=r_grand_tot, column=5).fill = grand_total_fill
    cell_grand_tot.fill = grand_total_fill
    ws1.cell(row=r_grand_tot, column=5).border = double_bottom_border
    cell_grand_tot.border = double_bottom_border

    ws1.column_dimensions['A'].width = 12
    ws1.column_dimensions['B'].width = 75
    ws1.column_dimensions['C'].width = 10
    ws1.column_dimensions['D'].width = 16
    ws1.column_dimensions['E'].width = 20
    ws1.column_dimensions['F'].width = 22

    # -------------------------------------------------------------
    # HOJA 2: ESCANDALLO DE COSTES INTERNOS HAM
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Escandallo Costes HAM")
    internal_header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")

    ws2["A1"] = "HAM ESTRUCTURAS METÁLICAS - ESCANDALLO DE COMPRA Y COSTES DIRECTOS INTERNOS"
    ws2["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E1B4B")
    ws2["A2"] = "CONFIDENCIAL: USO INTERNO EXCLUSIVO · TARIFAS ACTUALIZADAS (01: 25€, 02/03: 135€, 04/05/06: 95€) · MEDIOS INTEGRADOS"
    ws2["A2"].font = Font(name="Segoe UI", size=9.5, bold=True, color="EF4444")

    ws2["A4"] = "PARÁMETROS DE COMPRA DIRECTA HAM:"
    ws2["A4"].font = bold_font
    ws2["A5"] = "Coste compra Panel Sándwich HDX55 2 caras (40 mm PIR):"
    ws2["B5"] = 36.00
    ws2["C5"] = "€/m² (acopio fábrica Península)"
    ws2["A6"] = "Coste compra Panel/Chapa HDX55 Voladizos 2 caras:"
    ws2["B6"] = 34.00
    ws2["C6"] = "€/m²"
    ws2["A7"] = "Coste compra Chapa grecada HDX55 falso techo / tabica:"
    ws2["B7"] = 20.00
    ws2["C7"] = "€/m²"
    ws2["A8"] = "Tarifa interna mano de obra oficial montaje:"
    ws2["B8"] = 28.00
    ws2["C8"] = "€/hora"

    for r in range(5, 9):
        ws2[f"A{r}"].font = bold_font
        ws2[f"B{r}"].font = bold_font
        ws2[f"C{r}"].font = regular_font
        ws2[f"B{r}"].number_format = '#,##0.00 €'

    # Tabla de Costes Directos
    ws2["A10"] = "DESGLOSE DE COSTES DIRECTOS DE EJECUCIÓN (CDT HAM)"
    ws2["A10"].font = bold_font

    cost_headers = ["Capítulo", "Concepto de Coste Interno", "Unidad", "Medición", "Coste Unit. HAM (€)", "Coste Total HAM (€)"]
    for c_idx, head in enumerate(cost_headers, 1):
        cell = ws2.cell(row=11, column=c_idx, value=head)
        cell.fill = internal_header_fill
        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center" if c_idx not in [2] else "left", vertical="center")

    cost_data = [
        ["MAT-01", "Panel sándwich 40 mm PIR con acabado HDX55 en 2 caras cubiertas (1.371 m² + 5% solapes/mermas)", "m²", 1440.00, "=B5"],
        ["MAT-02", "Panel / chapa grecada acabado HDX55 en 2 caras para voladizos sup. (533 m² + 5% mermas)", "m²", 560.00, "=B6"],
        ["MAT-03", "Chapa grecada acabado HDX55 en 2 caras para falso techo inferior voladizo (533,60 m² + 5%)", "m²", 560.00, "=B7"],
        ["MAT-04", "Chapa grecada acabado HDX55 en 2 caras para tabica vertical de fachada (186,50 m² + 5%)", "m²", 196.00, "=B7"],
        ["MAT-05", "Tornillería roscachapa de acero galvanizado con arandela EPDM, accesorios fijación y masilla butilo", "lote", 1.00, 6400.00],
        ["MAT-06", "Rematería troquelada y plegada a medida con revestimiento HDX55 (cumbreras, petos, baberos)", "lote", 1.00, 5800.00],
        ["MAT-07", "Imprimación epoxi anticorrosiva de alto contenido en sólidos para saneado viguetas voladizos", "lote", 1.00, 1450.00],
        ["MAT-08", "Materiales perfiles, peldaños y cancelas en ACERO INOXIDABLE para 2 escalas con crinolina", "lote", 1.00, 3700.00],
        ["MAT-09", "Materiales sistema línea de vida LuxTop® (cable AISI 316, postes absorbedores, tensor y carros)", "lote", 1.00, 10300.00],
        ["MO-01", "Mano de obra especializada montaje HAM obra base (cuadrilla de 4 oficiales ~1050 h)", "horas", 1050.00, "=B8"],
        ["MO-02", "Mano de obra taller y montaje 2 escalas inox + instalación/ensayo línea LuxTop (~100 h)", "horas", 100.00, "=B8"],
        ["LOG-01", "Fletes marítimos Península-Menorca para transporte de trailers especiales paneles gran formato", "lote", 1.00, 5600.00],
        ["LOG-02", "Camión grúa pluma autopropulsada para descarga e izado distribuido sobre pórticos", "lote", 1.00, 4600.00],
        ["LOG-03", "Alquiler de 2 plataformas elevadoras móviles de personal (PEMP) articuladas para montaje", "lote", 1.00, 4800.00],
        ["PRL-01", "Protecciones colectivas: redes anticaídas bajo faldones UNE-EN 1263-1 y líneas de vida EN 795", "lote", 1.00, 3800.00],
        ["PRL-02", "Canon de vertedero insular y tasas de centro de reciclaje autorizado de Menorca", "lote", 1.00, 2400.00]
    ]

    for idx, c_row in enumerate(cost_data, start=12):
        ws2.cell(row=idx, column=1, value=c_row[0]).alignment = Alignment(horizontal="center", vertical="top")
        ws2.cell(row=idx, column=2, value=c_row[1]).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws2.cell(row=idx, column=3, value=c_row[2]).alignment = Alignment(horizontal="center", vertical="top")
        ws2.cell(row=idx, column=4, value=c_row[3]).alignment = Alignment(horizontal="right", vertical="top")
        ws2.cell(row=idx, column=5, value=c_row[4]).alignment = Alignment(horizontal="right", vertical="top")
        ws2.cell(row=idx, column=6, value=f"=D{idx}*E{idx}").alignment = Alignment(horizontal="right", vertical="top")

        for col_i in range(1, 7):
            cell = ws2.cell(row=idx, column=col_i)
            cell.font = regular_font
            cell.border = thin_border
            if col_i == 4:
                cell.number_format = '#,##0.00'
            elif col_i in [5, 6]:
                cell.number_format = '#,##0.00 €'

    # Fila Coste Directo Total
    r_cdt = 12 + len(cost_data)
    ws2.cell(row=r_cdt, column=5, value="COSTE DIRECTO TOTAL GLOBAL (CDT HAM):").font = bold_font
    ws2.cell(row=r_cdt, column=5).alignment = Alignment(horizontal="right")
    cell_cdt = ws2.cell(row=r_cdt, column=6, value=f"=SUM(F12:F{r_cdt-1})")
    cell_cdt.font = bold_font
    cell_cdt.number_format = '#,##0.00 €'
    cell_cdt.fill = section_fill
    ws2.cell(row=r_cdt, column=5).fill = section_fill

    # Cuadro de Comparativa y Rentabilidad
    r_rent = r_cdt + 2
    ws2.cell(row=r_rent, column=1, value="CUADRO DE RENTABILIDAD Y BENEFICIO COMERCIAL").font = bold_font

    headers_comp = ["Concepto Económico", "Obra Base (€)", "Opcionales (€)", "Total Global (€)", "% Margen Global"]
    for col_i, hc in enumerate(headers_comp, 1):
        c = ws2.cell(row=r_rent+1, column=col_i, value=hc)
        c.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        c.font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")

    comp_rows = [
        ["Precio de Venta al Cliente (Base Imp.):", 394382.00, 30740.00, "=B32+C32", "100.00%"],
        ["Coste Directo Total HAM:", 150250.00, 16800.00, "=B33+C33", "=D33/D32"],
        ["Margen Bruto de Explotación HAM:", "=B32-B33", "=C32-C33", "=D32-D33", "=D34/D32"]
    ]

    for c_idx, c_vals in enumerate(comp_rows, start=r_rent+2):
        ws2.cell(row=c_idx, column=1, value=c_vals[0]).font = bold_font
        for col_j, val in enumerate(c_vals[1:], start=2):
            c_cell = ws2.cell(row=c_idx, column=col_j, value=val)
            c_cell.font = bold_font
            c_cell.alignment = Alignment(horizontal="right")
            if col_j in [2, 3, 4]:
                c_cell.number_format = '#,##0.00 €'
            elif col_j == 5:
                c_cell.number_format = '0.00%'

        for col_i in range(1, 6):
            ws2.cell(row=c_idx, column=col_i).border = thin_border
            if c_idx == r_rent+4:
                ws2.cell(row=c_idx, column=col_i).fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

    ws2.column_dimensions['A'].width = 44
    ws2.column_dimensions['B'].width = 20
    ws2.column_dimensions['C'].width = 20
    ws2.column_dimensions['D'].width = 20
    ws2.column_dimensions['E'].width = 18
    ws2.column_dimensions['F'].width = 22

    xlsx_path = os.path.join(os.path.dirname(__file__), "presupuesto.xlsx")
    wb.save(xlsx_path)
    print(f"Excel guardado exitosamente con 2 hojas en: {xlsx_path}")

    # Generar versión exclusiva para cliente sin la hoja de escandallo interno
    wb_cliente = load_workbook(xlsx_path)
    if "Escandallo Costes HAM" in wb_cliente.sheetnames:
        wb_cliente.remove(wb_cliente["Escandallo Costes HAM"])
        xlsx_cliente_path = os.path.join(os.path.dirname(__file__), "presupuesto-cliente.xlsx")
        wb_cliente.save(xlsx_cliente_path)
        print(f"Excel cliente guardado exitosamente (sin costes) en: {xlsx_cliente_path}")

def generate_html_files():
    data = load_data()
    dir_path = os.path.dirname(__file__)

    # 1. HTML CLIENTE (3 PÁGINAS LIMPIAS Y ELEGANTES)
    html_cliente = f"""<!doctype html>
<html lang="es">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Presupuesto {data['numeroPresupuesto']}: Sustitución Cubierta y Voladizos Terminal de Carga - Aeropuerto de Menorca</title>
  <style>
    @page {{ 
      size: A4 portrait; 
      margin: 6mm 6mm 6mm 6mm;
    }}
    * {{ box-sizing: border-box; }}
    body {{ 
      margin: 0; 
      background: #eaedf0; 
      color: #1e262c; 
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; 
      font-size: 9.8px; 
      line-height: 1.28; 
    }}
    .page-container {{
      width: 210mm;
      margin: 0 auto;
      background: #ffffff;
    }}
    .page {{ 
      padding: 6.5mm 8.5mm 5.5mm;
      min-height: 285mm;
      max-height: 286mm;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      page-break-after: always;
      position: relative;
      box-sizing: border-box;
    }}
    .page:last-child {{
      page-break-after: avoid;
    }}

    header {{ 
      background: #0f172a; 
      color: #fff; 
      padding: 7px 13px; 
      border-radius: 6px;
      display: grid; 
      grid-template-columns: 1fr auto; 
      gap: 12px; 
      align-items: center; 
      border-bottom: 3.5px solid #b91c1c;
      margin-bottom: 6px;
    }}
    .brand {{ display: flex; align-items: center; gap: 10px; }}
    .logo {{ width: 38px; height: 38px; border: 1px solid rgba(255,255,255,.2); border-radius: 6px; padding: 3px; background: rgba(255,255,255,0.08); }}
    h1 {{ margin: 0; font-size: 16px; font-weight: 700; letter-spacing: -.02em; color: #ffffff; }}
    .subtitle {{ margin: 1px 0 0; color: #94a3b8; font-size: 9.5px; font-weight: 500; }}
    .meta {{ text-align: right; color: #cbd5e1; font-size: 8.8px; line-height: 1.35; }}
    .meta strong {{ color: #ffffff; display: block; font-size: 11px; margin-bottom: 2px; }}

    .client-card {{
      display: grid;
      grid-template-columns: 1.25fr 1fr;
      gap: 8px;
      background: #f8fafc;
      border: 1px solid #e2e8f0;
      border-radius: 6px;
      padding: 6px 10px;
      margin-bottom: 7px;
    }}
    .client-item h4 {{ margin: 0 0 1px; font-size: 8.5px; text-transform: uppercase; color: #64748b; letter-spacing: .05em; }}
    .client-item p {{ margin: 0; font-size: 10.2px; font-weight: 600; color: #0f172a; }}
    .client-item span {{ font-size: 9.2px; color: #475569; font-weight: normal; display: block; }}

    .kpi-grid {{
      display: grid;
      grid-template-columns: repeat(4, 1fr);
      gap: 7px;
      margin-bottom: 7px;
    }}
    .kpi-card {{
      background: #f1f5f9;
      border: 1px solid #cbd5e1;
      border-left: 4px solid #b91c1c;
      border-radius: 5px;
      padding: 5px 7px;
    }}
    .kpi-title {{ font-size: 8px; text-transform: uppercase; color: #64748b; font-weight: 700; margin-bottom: 1px; }}
    .kpi-val {{ font-size: 13px; font-weight: 800; color: #0f172a; line-height: 1.15; }}
    .kpi-sub {{ font-size: 8px; color: #475569; margin-top: 1px; }}

    h2 {{ 
      margin: 6px 0 3px; 
      font-size: 9.8px; 
      text-transform: uppercase; 
      letter-spacing: .05em; 
      color: #0f172a; 
      border-bottom: 1.5px solid #e2e8f0;
      padding-bottom: 2px;
      display: flex;
      justify-content: space-between;
      align-items: center;
    }}
    h2 span {{ font-size: 8px; color: #64748b; text-transform: none; font-weight: normal; }}

    .solution-box {{
      background: #fff;
      border: 1px solid #e2e8f0;
      border-radius: 6px;
      padding: 6px 8px;
      margin-bottom: 6px;
      font-size: 9.1px;
      line-height: 1.33;
      color: #334155;
    }}
    .solution-box ul {{
      margin: 2px 0 0;
      padding-left: 13px;
    }}
    .solution-box li {{
      margin-bottom: 1.5px;
    }}

    .remates-box {{
      background: #fdfefe;
      border: 1px solid #cbd5e1;
      border-left: 3.5px solid #0284c7;
      border-radius: 5px;
      padding: 5px 8px;
      margin-bottom: 6px;
      font-size: 8.8px;
      line-height: 1.3;
    }}
    .remates-title {{
      font-weight: 700;
      color: #0369a1;
      margin-bottom: 3px;
      text-transform: uppercase;
      font-size: 8.8px;
    }}
    .remates-grid {{
      display: grid;
      grid-template-columns: repeat(2, 1fr);
      gap: 4px 10px;
    }}
    .remates-item {{
      display: flex;
      justify-content: space-between;
      border-bottom: 1px dashed #e2e8f0;
      padding-bottom: 1px;
    }}

    table.data-table {{
      width: 100%;
      border-collapse: collapse;
      margin-bottom: 5px;
      font-size: 8.9px;
    }}
    table.data-table th {{
      background: #1e293b;
      color: #fff;
      padding: 4px 6px;
      text-align: left;
      font-weight: 600;
      font-size: 8.6px;
    }}
    table.data-table.optional th {{
      background: #1e3a8a;
    }}
    table.data-table td {{
      padding: 3.5px 6px;
      border-bottom: 1px solid #e2e8f0;
      vertical-align: top;
    }}
    table.data-table tr:nth-child(even) {{
      background: #f8fafc;
    }}
    .td-right {{ text-align: right; white-space: nowrap; }}
    .td-center {{ text-align: center; }}

    .partida-title {{ font-weight: 700; color: #0f172a; font-size: 9px; line-height: 1.25; }}
    .partida-desc {{ font-size: 8.2px; color: #475569; margin-top: 1.5px; line-height: 1.22; text-align: justify; }}
    .partida-code {{ background: #e2e8f0; color: #0f172a; padding: 1px 3px; border-radius: 3px; font-weight: 700; font-size: 8px; }}
    .partida-code-opc {{ background: #dbeafe; color: #1e40af; padding: 1px 3px; border-radius: 3px; font-weight: 700; font-size: 8px; }}

    .summary-cards-container {{
      display: grid;
      grid-template-columns: 1fr 1fr 1.3fr;
      gap: 8px;
      margin: 8px 0;
    }}
    .sum-box {{
      background: #f8fafc;
      border: 1px solid #cbd5e1;
      border-radius: 6px;
      padding: 7px 9px;
      text-align: center;
    }}
    .sum-box.highlight {{
      background: #ecfdf5;
      border: 1.5px solid #10b981;
    }}
    .sum-box.optional-box {{
      background: #eff6ff;
      border: 1.5px solid #3b82f6;
    }}
    .sum-box h6 {{
      margin: 0 0 2px;
      font-size: 8.5px;
      text-transform: uppercase;
      color: #64748b;
    }}
    .sum-val {{
      font-size: 15px;
      font-weight: 800;
      color: #0f172a;
    }}
    .sum-sub {{
      font-size: 8.8px;
      color: #475569;
      margin-top: 2px;
    }}

    .conditions-grid {{
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 8px;
      background: #f8fafc;
      border: 1px solid #e2e8f0;
      border-radius: 6px;
      padding: 6px 10px;
      margin-bottom: 7px;
    }}
    .cond-item h5 {{
      margin: 0 0 2px;
      font-size: 8.5px;
      text-transform: uppercase;
      color: #0f172a;
      border-bottom: 1px solid #cbd5e1;
      padding-bottom: 1px;
    }}
    .cond-item p {{
      margin: 0;
      font-size: 8.4px;
      color: #334155;
      line-height: 1.28;
    }}

    .signatures {{
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 12px;
      margin-top: 6px;
    }}
    .sign-box {{
      text-align: center;
      border: 1px dashed #cbd5e1;
      border-radius: 5px;
      padding: 6px;
      min-height: 48px;
      font-size: 8.3px;
      color: #64748b;
    }}

    footer {{ 
      border-top: 1px solid #cbd5e1; 
      padding-top: 2.5px; 
      display: flex; 
      justify-content: space-between; 
      font-size: 7.8px; 
      color: #64748b; 
    }}
    .badge-hdx {{
      background: #fee2e2;
      color: #991b1b;
      padding: 1px 4px;
      border-radius: 3px;
      font-weight: 700;
      font-size: 8px;
      display: inline-block;
    }}
    .badge-inox {{
      background: #dbeafe;
      color: #1e40af;
      padding: 1px 4px;
      border-radius: 3px;
      font-weight: 700;
      font-size: 8px;
      display: inline-block;
    }}
    .badge-luxtop {{
      background: #fef3c7;
      color: #92400e;
      padding: 1px 4px;
      border-radius: 3px;
      font-weight: 700;
      font-size: 8px;
      display: inline-block;
    }}
  </style>
</head>
<body>
<div class="page-container">

  <!-- ==================== PÁGINA 1: MEMORIA, DESGLOSE SUPERFICIES Y REMATERÍA ==================== -->
  <div class="page">
    <div>
      <header>
        <div class="brand">
          <img src="logo-ham.svg" class="logo" alt="HAM Logo">
          <div>
            <h1>HAM ESTRUCTURAS METÁLICAS</h1>
            <div class="subtitle">DIXIS RECURSOS Y SOLUCIONES, SL · Licitaciones y Envolventes Industriales</div>
          </div>
        </div>
        <div class="meta">
          <strong>OFERTA TÉCNICA {data['numeroPresupuesto']} (REV5)</strong>
          Ref: {data['referencia']}<br>
          Fecha: 28 de octubre de 2026<br>
          Validez: 30 días naturales
        </div>
      </header>

      <div class="client-card">
        <div class="client-item">
          <h4>Cliente y Proyectista</h4>
          <p>{data['cliente']['empresa']}</p>
          <span>Atn: {data['cliente']['atencion']} | Tlf: {data['cliente']['telefono']}</span>
          <span>Web: {data['cliente']['web']}</span>
        </div>
        <div class="client-item">
          <h4>Proyecto y Emplazamiento</h4>
          <p>{data['cliente']['proyecto']}</p>
          <span>Promotor: {data['cliente']['promotorFinal']}</span>
          <span>Ubicación: {data['cliente']['emplazamiento']}, {data['cliente']['poblacion']} ({data['cliente']['isla']})</span>
        </div>
      </div>

      <div class="kpi-grid">
        <div class="kpi-card">
          <div class="kpi-title">Superficie Total</div>
          <div class="kpi-val">2.624,10 m²</div>
          <div class="kpi-sub">Cubierta (1.371 m²) + Voladizos y Tabica (1.253,10 m²)</div>
        </div>
        <div class="kpi-card">
          <div class="kpi-title">Sistema de Cubierta</div>
          <div class="kpi-val">Panel PIR 40mm</div>
          <div class="kpi-sub">3 grecas con tapajuntas (B-s1,d0)</div>
        </div>
        <div class="kpi-card">
          <div class="kpi-title">Acabado Protección</div>
          <div class="kpi-val">Granite® HDX55</div>
          <div class="kpi-sub">55 µm en AMBAS CARAS (C5-M marino)</div>
        </div>
        <div class="kpi-card">
          <div class="kpi-title">Fijación Estanca</div>
          <div class="kpi-val">Galv. + EPDM</div>
          <div class="kpi-sub">Tornillería roscachapa con junta EPDM</div>
        </div>
      </div>

      <h2>1. MEMORIA TÉCNICA Y SISTEMA PROPUESTO <span>Protección Marina Severa C5-M, Medios Integrados y Seguridad</span></h2>
      <div class="solution-box">
        Atendiendo a las especificaciones y ampliaciones requeridas por <strong>UG21 Engineering</strong> para la terminal de carga del Aeropuerto de Menorca, se formula propuesta técnica de renovación integral que comprende la sustitución de la cubierta principal, la cobertura superior e inferior de los voladizos de muelles y la tabica vertical de fachada:
        <ul>
          <li><strong>Recubrimiento Granite® HDX55 (55 micras) en AMBAS CARAS:</strong> Sistema orgánico multicapa sobre chapa estructural galvanizada Z275. Máxima categoría de resistencia a la corrosión <strong>RC5 / C5-M (Marino Severo)</strong> y radiación UV <strong>RUV4</strong> (UNE-EN 10169). Inmune a la niebla salina marina de Menorca y gases de queroseno de turbinas aéreas. Aplicado al 100% en cubierta, caras superior e inferior de voladizos y tabica.</li>
          <li><strong>Medios de Elevación y Fletes Íntegramente Repercutidos:</strong> Todos los transportes marítimos Península-Menorca en trailers especiales, camión grúa telescópica autopropulsada para izado distribuido en pórticos y plataformas elevadoras PEMP se encuentran totalmente repercutidos e integrados en los precios unitarios de cada partida.</li>
          <li><strong>Panel Sándwich PIR 40 mm:</strong> Espesor 40 mm, densidad 40 kg/m³, transmitancia reducida (U = 0,51 W/m²·K) y clasificación al fuego <strong>B-s1,d0 / Broof(t1)</strong> según UNE-EN 13501-1.</li>
          <li><strong>Tornillería Roscachapa Galvanizada con EPDM:</strong> Fijaciones mecánicas autotaladrantes de acero galvanizado con arandela vulcanizada de estanqueidad en EPDM y cordón hidro-expansivo continuo en solapes.</li>
          <li><strong>Escaleras Inox con Crinolina (Opcionales):</strong> Conforme a <strong>UNE-EN ISO 14122-4</strong> íntegramente en acero inoxidable AISI 304/316, jaula dorsal, cancela batiente con muelle y portón inferior de bloqueo con candado (protocolo AENA).</li>
          <li><strong>Línea de Vida Homologada LUXTOP® (Opcional):</strong> Cable de acero inoxidable AISI 316 continuo (Ø 8 mm) conforme a <strong>UNE-EN 795:2012 Tipo C y CEN/TS 16415</strong> para hasta 4 operarios simultáneos, postes sobre faldón con absorbedores dinámicos y carros pasantes continuos sin desconexión.</li>
        </ul>
      </div>

      <h2>2. DESGLOSE DE REMATERÍA TÉCNICA HDX55 (PARTIDA 07) <span>Metrado verificado según Planimetrías (489,15 ml + 10 ud singulares)</span></h2>
      <div class="remates-box">
        <div class="remates-title">Detalle de mediciones lineales para control y supervisión técnica:</div>
        <div class="remates-grid">
          <div class="remates-item"><span>• Cumbrera troquelada faldón a 2 aguas:</span><strong>50,23 ml</strong></div>
          <div class="remates-item"><span>• Testeros / hastiales laterales a 2 aguas (4 x 10,40m):</span><strong>41,60 ml</strong></div>
          <div class="remates-item"><span>• Testeros / hastiales faldón a 1 agua (2 x 10,87m):</span><strong>21,74 ml</strong></div>
          <div class="remates-item"><span>• Coronación babero contra peto superior (1 agua):</span><strong>30,22 ml</strong></div>
          <div class="remates-item"><span>• Remate de quiebro/salto entre faldones:</span><strong>10,40 ml</strong></div>
          <div class="remates-item"><span>• Goterones frontales voladizos Este y Oeste (80,44 x 2):</span><strong>160,88 ml</strong></div>
          <div class="remates-item"><span>• Cierres laterales de testero en voladizos (4 x 3,30m):</span><strong>13,20 ml</strong></div>
          <div class="remates-item"><span>• Babero posterior voladizo contra fachada (80,44 x 2):</span><strong>160,88 ml</strong></div>
        </div>
        <div style="margin-top: 3px; display: flex; justify-content: space-between; border-top: 1px solid #cbd5e1; padding-top: 2px;">
          <span><strong>Total Desarrollo Lineal Rematería HDX55: 489,15 ml</strong></span>
          <span><strong>Puntos singulares (cajeados chimeneas/bajantes): 10 ud</strong></span>
        </div>
      </div>

      <h2>3. DESGLOSE DE SUPERFICIES DE ACTUACIÓN <span>Terminal de Carga Aeropuerto de Menorca</span></h2>
      <table class="data-table">
        <thead>
          <tr>
            <th style="width: 40px;" class="td-center">Sector</th>
            <th>Denominación del Elemento</th>
            <th>Tipología Constructiva</th>
            <th class="td-right">Superficie</th>
            <th class="td-center">Acabado Dos Caras</th>
            <th class="td-center">Ambiente Normativo</th>
          </tr>
        </thead>
        <tbody>
          <tr>
            <td class="td-center"><strong>01</strong></td>
            <td><strong>Cubierta Principal a Dos Aguas</strong></td>
            <td>Panel Sándwich PIR 40 mm con tapajuntas</td>
            <td class="td-right"><strong>1.042,50 m²</strong></td>
            <td class="td-center"><span class="badge-hdx">HDX55 / HDX55</span></td>
            <td class="td-center">Marino C5-M / AENA</td>
          </tr>
          <tr>
            <td class="td-center"><strong>02</strong></td>
            <td><strong>Cubierta Faldón a Un Agua</strong></td>
            <td>Panel Sándwich PIR 40 mm con cortes y adaptación a petos</td>
            <td class="td-right"><strong>328,50 m²</strong></td>
            <td class="td-center"><span class="badge-hdx">HDX55 / HDX55</span></td>
            <td class="td-center">Marino C5-M / AENA</td>
          </tr>
          <tr>
            <td class="td-center"><strong>03</strong></td>
            <td><strong>Voladizos de Fachadas (Cobertura Superior)</strong></td>
            <td>Chapa/panel grecado sobre ménsulas metálicas saneadas</td>
            <td class="td-right"><strong>533,00 m²</strong></td>
            <td class="td-center"><span class="badge-hdx">HDX55 / HDX55</span></td>
            <td class="td-center">Marino C5-M / AENA</td>
          </tr>
          <tr>
            <td class="td-center"><strong>04</strong></td>
            <td><strong>Lado Inferior de Voladizos (Cielo Raso)</strong></td>
            <td>Chapa grecada fijada bajo viguetas en ménsula</td>
            <td class="td-right"><strong>533,60 m²</strong></td>
            <td class="td-center"><span class="badge-hdx">HDX55 / HDX55</span></td>
            <td class="td-center">Marino C5-M / AENA</td>
          </tr>
          <tr>
            <td class="td-center"><strong>05</strong></td>
            <td><strong>Tabica Vertical de Fachada sobre Marquesinas</strong></td>
            <td>Chapa grecada en paramento vertical de peto/fachada</td>
            <td class="td-right"><strong>186,50 m²</strong></td>
            <td class="td-center"><span class="badge-hdx">HDX55 / HDX55</span></td>
            <td class="td-center">Marino C5-M / AENA</td>
          </tr>
          <tr style="background: #f1f5f9; font-weight: 700;">
            <td colspan="3" style="text-align: right; text-transform: uppercase;">TOTAL SUPERFICIE DE ACTUACIÓN:</td>
            <td class="td-right" style="color: #b91c1c; font-size: 9.8px;">2.624,10 m²</td>
            <td colspan="2" class="td-center">100% Tratamiento Anticorrosión HDX55 en 2 caras</td>
          </tr>
        </tbody>
      </table>
    </div>

    <footer>
      <span>HAM Estructuras Metálicas · Pol. Ind. Sant Lluís, Menorca · Tel: 971 35 20 18</span>
      <span>Página 1 de 3 · Oferta P-2026-0044 (REV5)</span>
    </footer>
  </div>

  <!-- ==================== PÁGINA 2: CUADRO DE PARTIDAS BASE ==================== -->
  <div class="page">
    <div>
      <header>
        <div class="brand">
          <img src="logo-ham.svg" class="logo" alt="HAM Logo">
          <div>
            <h1>HAM ESTRUCTURAS METÁLICAS</h1>
            <div class="subtitle">VALORACIÓN ECONÓMICA DE PARTIDAS BASE (MEDIOS INTEGRADOS)</div>
          </div>
        </div>
        <div class="meta">
          <strong>PRESUPUESTO {data['numeroPresupuesto']} (REV5)</strong>
          Ref: {data['referencia']}<br>
          Terminal Carga AENA Menorca
        </div>
      </header>

      <h2>4. CUADRO DE VALORACIÓN ECONÓMICA DE PARTIDAS BASE <span>Obra Principal Cubiertas, Voladizos y Tabica (Medios Integrados)</span></h2>
      <table class="data-table">
        <thead>
          <tr>
            <th style="width: 25px;" class="td-center">Nº</th>
            <th>Concepto y Resumen Descriptivo</th>
            <th style="width: 30px;" class="td-center">Ud.</th>
            <th style="width: 55px;" class="td-right">Cantidad</th>
            <th style="width: 65px;" class="td-right">P. Unitario</th>
            <th style="width: 78px;" class="td-right">Importe (€)</th>
          </tr>
        </thead>
        <tbody>
"""

    for p in data["partidas"]:
        html_cliente += f"""          <tr>
            <td class="td-center"><span class="partida-code">{p['id']}</span></td>
            <td>
              <div class="partida-title">{p['concepto']}</div>
              <div class="partida-desc">{p['descripcion']}</div>
            </td>
            <td class="td-center">{p['unidad']}</td>
            <td class="td-right">{p['cantidad']:,.2f}</td>
            <td class="td-right">{p['precioUnitario']:,.2f} €</td>
            <td class="td-right"><strong>{p['importe']:,.2f} €</strong></td>
          </tr>
"""

    html_cliente += f"""          <tr style="background: #f1f5f9; font-weight: 700;">
            <td colspan="5" style="text-align: right; text-transform: uppercase;">SUBTOTAL BASE IMPONIBLE OBRA BASE:</td>
            <td class="td-right" style="color: #b91c1c; font-size: 9.8px;">{data['totalesEconomicos']['baseImponibleBase']:,.2f} €</td>
          </tr>
          <tr style="background: #ffffff;">
            <td colspan="5" style="text-align: right;">I.V.A. (21%) OBRA BASE:</td>
            <td class="td-right">{data['totalesEconomicos']['ivaBase']:,.2f} €</td>
          </tr>
          <tr style="background: #fef2f2; font-weight: 800; border-top: 1.5px solid #b91c1c;">
            <td colspan="5" style="text-align: right; text-transform: uppercase; color: #991b1b;">TOTAL OBRA BASE (IVA INCLUIDO):</td>
            <td class="td-right" style="color: #991b1b; font-size: 10.5px;">{data['totalesEconomicos']['totalBaseConIva']:,.2f} €</td>
          </tr>
        </tbody>
      </table>
    </div>

    <footer>
      <span>HAM Estructuras Metálicas · Menorca · info@hamenorca.com</span>
      <span>Página 2 de 3 · Oferta P-2026-0044 (REV5)</span>
    </footer>
  </div>

  <!-- ==================== PÁGINA 3: OPCIONALES, TOTALES Y CONDICIONES ==================== -->
  <div class="page">
    <div>
      <header>
        <div class="brand">
          <img src="logo-ham.svg" class="logo" alt="HAM Logo">
          <div>
            <h1>HAM ESTRUCTURAS METÁLICAS</h1>
            <div class="subtitle">PARTIDAS OPCIONALES, TOTALES Y CONDICIONES COMERCIALES</div>
          </div>
        </div>
        <div class="meta">
          <strong>PRESUPUESTO {data['numeroPresupuesto']} (REV5)</strong>
          Ref: {data['referencia']}<br>
          Terminal Carga AENA Menorca
        </div>
      </header>

      <h2>5. PARTIDAS OPCIONALES: SEGURIDAD EN ALTURA (ESCALAS INOX Y LÍNEA LUXTOP®) <span>Homologadas UNE-EN ISO 14122-4 / UNE-EN 795</span></h2>
      <table class="data-table optional">
        <thead>
          <tr>
            <th style="width: 48px;" class="td-center">Código</th>
            <th>Concepto y Elementos de Seguridad Integrados</th>
            <th style="width: 30px;" class="td-center">Ud.</th>
            <th style="width: 55px;" class="td-right">Cantidad</th>
            <th style="width: 65px;" class="td-right">P. Unitario</th>
            <th style="width: 78px;" class="td-right">Importe (€)</th>
          </tr>
        </thead>
        <tbody>
"""

    for op in data["partidasOpcionales"]:
        badge_class = "badge-luxtop" if "LUXTOP" in op['codigo'] else "badge-inox"
        badge_text = "LUXTOP® EN 795" if "LUXTOP" in op['codigo'] else "ACERO INOX"
        html_cliente += f"""          <tr>
            <td class="td-center"><span class="partida-code-opc">{op['id']}</span></td>
            <td>
              <div class="partida-title">{op['concepto']} <span class="{badge_class}">{badge_text}</span></div>
              <div class="partida-desc">{op['descripcion']}</div>
            </td>
            <td class="td-center">{op['unidad']}</td>
            <td class="td-right">{op['cantidad']:,.2f}</td>
            <td class="td-right">{op['precioUnitario']:,.2f} €</td>
            <td class="td-right"><strong>{op['importe']:,.2f} €</strong></td>
          </tr>
"""

    html_cliente += f"""          <tr style="background: #eff6ff; font-weight: 700;">
            <td colspan="5" style="text-align: right; text-transform: uppercase;">SUBTOTAL BASE IMPONIBLE OPCIONALES:</td>
            <td class="td-right" style="color: #1e3a8a; font-size: 9.5px;">{data['totalesEconomicos']['baseImponibleOpcionales']:,.2f} €</td>
          </tr>
          <tr style="background: #ffffff;">
            <td colspan="5" style="text-align: right;">I.V.A. (21%) OPCIONALES:</td>
            <td class="td-right">{data['totalesEconomicos']['ivaOpcionales']:,.2f} €</td>
          </tr>
          <tr style="background: #f0fdf4; font-weight: 800; border-top: 1.5px solid #10b981;">
            <td colspan="5" style="text-align: right; text-transform: uppercase; color: #065f46;">TOTAL OPCIONALES (IVA INCLUIDO):</td>
            <td class="td-right" style="color: #065f46; font-size: 10px;">{data['totalesEconomicos']['totalOpcionalesConIva']:,.2f} €</td>
          </tr>
        </tbody>
      </table>

      <div class="summary-cards-container">
        <div class="sum-box">
          <h6>Obra Base (Medios Integrados)</h6>
          <div class="sum-val">{data['totalesEconomicos']['baseImponibleBase']:,.2f} €</div>
          <div class="sum-sub">+ IVA 21%: {data['totalesEconomicos']['totalBaseConIva']:,.2f} €</div>
        </div>
        <div class="sum-box optional-box">
          <h6>Opcionales (2 Escalas Inox + LuxTop)</h6>
          <div class="sum-val" style="color: #1e3a8a;">+{data['totalesEconomicos']['baseImponibleOpcionales']:,.2f} €</div>
          <div class="sum-sub">+ IVA 21%: {data['totalesEconomicos']['totalOpcionalesConIva']:,.2f} €</div>
        </div>
        <div class="sum-box highlight">
          <h6>Total Conjunto (Base + Opcionales)</h6>
          <div class="sum-val" style="color: #047857;">{data['totalesEconomicos']['baseImponibleGeneral']:,.2f} €</div>
          <div class="sum-sub" style="color: #065f46; font-weight: 700;">TOTAL CON IVA: {data['totalesEconomicos']['totalGeneralConIva']:,.2f} €</div>
        </div>
      </div>

      <h2>6. CONDICIONES COMERCIALES, LOGÍSTICAS Y GARANTÍAS</h2>
      <div class="conditions-grid">
        <div class="cond-item">
          <h5>Forma de Pago y Contratación</h5>
          <p>• <strong>Forma de pago:</strong> A convenir con el cliente según avance y certificaciones de obra.<br>
          • <strong>Partidas opcionales:</strong> Facturación independiente según previa aceptación y aprobación técnica de la Dirección de Obra.</p>
        </div>
        <div class="cond-item">
          <h5>Plazos, AENA y Garantías</h5>
          <p>• <strong>Plazo de ejecución:</strong> {data['condicionesComerciales']['plazoEjecucion']}.<br>
          • <strong>Garantías:</strong> 10 años en Granite® HDX55, inox y LuxTop®; 5 años estanqueidad; 2 años montaje.<br>
          • <strong>Seguridad:</strong> Cumplimiento estricto UNE-EN ISO 14122-4, UNE-EN 795 y normativa de seguridad operacional AENA.</p>
        </div>
      </div>

      <div class="signatures">
        <div class="sign-box">
          <strong>Por la Empresa Contratista:</strong><br>
          HAM ESTRUCTURAS METÁLICAS<br>
          (Dixis Recursos y Soluciones, SL)<br>
          <em>Dpto. Técnico y Dirección de Obra</em>
        </div>
        <div class="sign-box">
          <strong>Conforme el Cliente / Dirección Técnica:</strong><br>
          CONSULTORES DE INGENIERÍA UG21<br>
          (UTE BTG-UG21 / AENA)<br>
          <em>Dña. Marina Gallego Fernández (Arquitecto)</em>
        </div>
      </div>
    </div>

    <footer>
      <span>HAM Estructuras Metálicas · Menorca · info@hamenorca.com</span>
      <span>Página 3 de 3 · Oferta P-2026-0044 (REV5)</span>
    </footer>
  </div>

</div>
</body>
</html>
"""

    with open(os.path.join(dir_path, "presupuesto-cliente.html"), "w", encoding="utf-8") as f:
        f.write(html_cliente)
    print("presupuesto-cliente.html actualizado con éxito.")

    # 2. HTML INTERNO TÉCNICO
    html_interno = f"""<!doctype html>
<html lang="es">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Dossier Interno Técnico {data['numeroPresupuesto']} - Costes HAM y Rentabilidad</title>
  <style>
    @page {{ 
      size: A4 portrait; 
      margin: 6mm 6mm 6mm 6mm;
    }}
    * {{ box-sizing: border-box; }}
    body {{ 
      margin: 0; 
      background: #f1f5f9; 
      color: #0f172a; 
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; 
      font-size: 9.8px; 
      line-height: 1.28; 
    }}
    .page-container {{
      width: 210mm;
      margin: 0 auto;
      background: #ffffff;
    }}
    .page {{ 
      padding: 6.5mm 8.5mm 5.5mm;
      min-height: 285mm;
      max-height: 286mm;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      page-break-after: always;
      position: relative;
      box-sizing: border-box;
    }}
    .page:last-child {{
      page-break-after: avoid;
    }}

    header {{ 
      background: #1e1b4b; 
      color: #fff; 
      padding: 7px 12px; 
      border-radius: 5px;
      display: grid; 
      grid-template-columns: 1fr auto; 
      gap: 10px; 
      align-items: center; 
      border-bottom: 3.5px solid #4f46e5;
      margin-bottom: 6px;
    }}
    .brand {{ display: flex; align-items: center; gap: 10px; }}
    .logo {{ width: 36px; height: 36px; border: 1px solid rgba(255,255,255,.2); border-radius: 5px; padding: 3px; background: rgba(255,255,255,0.08); }}
    h1 {{ margin: 0; font-size: 15.5px; font-weight: 700; color: #ffffff; }}
    .subtitle {{ margin: 1px 0 0; color: #a5b4fc; font-size: 9px; font-weight: 500; }}
    .meta {{ text-align: right; color: #c7d2fe; font-size: 8.5px; }}
    .meta strong {{ color: #ffffff; display: block; font-size: 10.5px; margin-bottom: 1px; }}

    h2 {{ 
      margin: 5px 0 3px; 
      font-size: 9.5px; 
      text-transform: uppercase; 
      letter-spacing: .05em; 
      color: #1e1b4b; 
      border-bottom: 1.5px solid #e2e8f0;
      padding-bottom: 2px;
    }}

    table.data-table {{
      width: 100%;
      border-collapse: collapse;
      margin-bottom: 5px;
      font-size: 8.8px;
    }}
    table.data-table th {{
      background: #1e1b4b;
      color: #fff;
      padding: 3.5px 5px;
      text-align: left;
      font-weight: 600;
      font-size: 8.5px;
    }}
    table.data-table td {{
      padding: 3px 5px;
      border-bottom: 1px solid #e2e8f0;
      vertical-align: top;
    }}
    table.data-table tr:nth-child(even) {{
      background: #f8fafc;
    }}
    .td-right {{ text-align: right; white-space: nowrap; }}
    .td-center {{ text-align: center; }}

    .margin-card {{
      background: #ecfdf5;
      border: 1.5px solid #10b981;
      border-radius: 6px;
      padding: 6px 10px;
      margin: 6px 0;
    }}
    .margin-grid {{
      display: grid;
      grid-template-columns: repeat(4, 1fr);
      gap: 8px;
      text-align: center;
    }}
    .margin-label {{ font-size: 8px; text-transform: uppercase; color: #065f46; font-weight: 700; }}
    .margin-val {{ font-size: 13.5px; font-weight: 800; color: #047857; line-height: 1.2; }}

    footer {{ 
      border-top: 1px solid #cbd5e1; 
      padding-top: 2px; 
      display: flex; 
      justify-content: space-between; 
      font-size: 7.5px; 
      color: #64748b; 
    }}
  </style>
</head>
<body>
<div class="page-container">

  <!-- ==================== PÁGINA 1: RESUMEN VENTAS Y DESGLOSE PARTIDAS BASE ==================== -->
  <div class="page">
    <div>
      <header>
        <div class="brand">
          <img src="logo-ham.svg" class="logo" alt="HAM Logo">
          <div>
            <h1>HAM ESTRUCTURAS METÁLICAS · CONTROL INTERNO</h1>
            <div class="subtitle">ESCANDALLO TÉCNICO, COSTES DIRECTOS Y RENTABILIDAD HAM (REV5)</div>
          </div>
        </div>
        <div class="meta">
          <strong>EXPEDIENTE {data['numeroPresupuesto']}</strong>
          Confidencial HAM · Uso Interno<br>
          Superficie Total: 2.624,10 m²
        </div>
      </header>

      <h2>1. CUADRO DE PRECIOS DE VENTA (TARIFAS 01: 25€, 02/03: 135€, 04/05/06: 95€)</h2>
      <table class="data-table">
        <thead>
          <tr>
            <th style="width: 25px;" class="td-center">Nº</th>
            <th>Concepto de Venta al Cliente</th>
            <th style="width: 30px;" class="td-center">Ud.</th>
            <th style="width: 55px;" class="td-right">Cantidad</th>
            <th style="width: 65px;" class="td-right">PVP Unitario</th>
            <th style="width: 78px;" class="td-right">Importe Venta</th>
          </tr>
        </thead>
        <tbody>
"""

    for p in data["partidas"]:
        html_interno += f"""          <tr>
            <td class="td-center"><strong>{p['id']}</strong></td>
            <td><strong>{p['concepto']}</strong></td>
            <td class="td-center">{p['unidad']}</td>
            <td class="td-right">{p['cantidad']:,.2f}</td>
            <td class="td-right">{p['precioUnitario']:,.2f} €</td>
            <td class="td-right"><strong>{p['importe']:,.2f} €</strong></td>
          </tr>
"""

    html_interno += f"""          <tr style="background: #f1f5f9; font-weight: 700;">
            <td colspan="5" style="text-align: right; text-transform: uppercase;">VENTA TOTAL OBRA BASE (BI):</td>
            <td class="td-right" style="color: #b91c1c;">{data['totalesEconomicos']['baseImponibleBase']:,.2f} €</td>
          </tr>
        </tbody>
      </table>

      <h2>2. PARTIDAS OPCIONALES DE SEGURIDAD (PRECIOS DE VENTA AL CLIENTE)</h2>
      <table class="data-table">
        <thead>
          <tr>
            <th style="width: 48px;" class="td-center">Código</th>
            <th>Concepto Opcional</th>
            <th style="width: 30px;" class="td-center">Ud.</th>
            <th style="width: 55px;" class="td-right">Cantidad</th>
            <th style="width: 65px;" class="td-right">PVP Unitario</th>
            <th style="width: 78px;" class="td-right">Importe Venta</th>
          </tr>
        </thead>
        <tbody>
"""

    for op in data["partidasOpcionales"]:
        html_interno += f"""          <tr>
            <td class="td-center"><strong>{op['id']}</strong></td>
            <td><strong>{op['concepto']}</strong></td>
            <td class="td-center">{op['unidad']}</td>
            <td class="td-right">{op['cantidad']:,.2f}</td>
            <td class="td-right">{op['precioUnitario']:,.2f} €</td>
            <td class="td-right"><strong>{op['importe']:,.2f} €</strong></td>
          </tr>
"""

    html_interno += f"""          <tr style="background: #eff6ff; font-weight: 700;">
            <td colspan="5" style="text-align: right; text-transform: uppercase;">VENTA TOTAL OPCIONALES (BI):</td>
            <td class="td-right" style="color: #1e3a8a;">{data['totalesEconomicos']['baseImponibleOpcionales']:,.2f} €</td>
          </tr>
          <tr style="background: #ecfdf5; font-weight: 800;">
            <td colspan="5" style="text-align: right; text-transform: uppercase; color: #065f46;">VENTA GLOBAL CONJUNTA (BASE + OPCIONALES):</td>
            <td class="td-right" style="color: #065f46; font-size: 10px;">{data['totalesEconomicos']['baseImponibleGeneral']:,.2f} €</td>
          </tr>
        </tbody>
      </table>
    </div>

    <footer>
      <span>HAM Estructuras Metálicas · Control de Gestión y Costes</span>
      <span>Dossier Técnico Interno Confidencial · P-2026-0044 (REV5) · Página 1 de 2</span>
    </footer>
  </div>

  <!-- ==================== PÁGINA 2: ESCANDALLO DETALLADO Y MARGEN DE EXPLOTACIÓN ==================== -->
  <div class="page">
    <div>
      <header>
        <div class="brand">
          <img src="logo-ham.svg" class="logo" alt="HAM Logo">
          <div>
            <h1>HAM ESTRUCTURAS METÁLICAS · ESCANDALLO INTERNO</h1>
            <div class="subtitle">DESGLOSE DE MATERIALES, MANO DE OBRA, LOGÍSTICA Y RENTABILIDAD</div>
          </div>
        </div>
        <div class="meta">
          <strong>EXPEDIENTE {data['numeroPresupuesto']}</strong>
          Confidencial HAM · Uso Interno<br>
          Coste Directo: {data['escandalloCostesInternosHAM']['costeDirectoTotalGlobal']:,.2f} €
        </div>
      </header>

      <h2>3. ESCANDALLO DE COSTES DIRECTOS DE EJECUCIÓN (CDT HAM)</h2>
      <table class="data-table">
        <thead>
          <tr>
            <th style="width: 50px;" class="td-center">Capítulo</th>
            <th>Concepto de Coste Interno</th>
            <th style="width: 30px;" class="td-center">Ud.</th>
            <th style="width: 50px;" class="td-right">Medición</th>
            <th style="width: 65px;" class="td-right">Coste Unit.</th>
            <th style="width: 78px;" class="td-right">Coste Total HAM</th>
          </tr>
        </thead>
        <tbody>
          <tr>
            <td class="td-center"><strong>MAT-01</strong></td>
            <td>Panel sándwich 40 mm PIR acabado HDX55 en 2 caras cubiertas (1.371 m² + 5% solapes/mermas)</td>
            <td class="td-center">m²</td>
            <td class="td-right">1.440,00</td>
            <td class="td-right">36,00 €</td>
            <td class="td-right"><strong>51.840,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>MAT-02</strong></td>
            <td>Panel / chapa grecada acabado HDX55 2 caras voladizos cara superior (533 m² + 5% mermas)</td>
            <td class="td-center">m²</td>
            <td class="td-right">560,00</td>
            <td class="td-right">34,00 €</td>
            <td class="td-right"><strong>19.040,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>MAT-03</strong></td>
            <td>Chapa grecada acabado HDX55 2 caras para falso techo inferior voladizo (533,60 m² + 5%)</td>
            <td class="td-center">m²</td>
            <td class="td-right">560,00</td>
            <td class="td-right">20,00 €</td>
            <td class="td-right"><strong>11.200,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>MAT-04</strong></td>
            <td>Chapa grecada acabado HDX55 2 caras para tabica vertical de fachada (186,50 m² + 5%)</td>
            <td class="td-center">m²</td>
            <td class="td-right">196,00</td>
            <td class="td-right">20,00 €</td>
            <td class="td-right"><strong>3.920,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>MAT-05</strong></td>
            <td>Tornillería roscachapa de acero galvanizado con arandela vulcanizada EPDM y cordón butilo</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">6.400,00 €</td>
            <td class="td-right"><strong>6.400,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>MAT-06</strong></td>
            <td>Rematería troquelada y plegada a medida con revestimiento HDX55 (489,15 ml + 10 ud singulares)</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">5.800,00 €</td>
            <td class="td-right"><strong>5.800,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>MAT-07</strong></td>
            <td>Imprimación epoxi anticorrosiva alto contenido en sólidos para saneado de viguetas voladizos</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">1.450,00 €</td>
            <td class="td-right"><strong>1.450,00 €</strong></td>
          </tr>
          <tr style="background: #eff6ff;">
            <td class="td-center"><strong style="color: #1e40af;">MAT-08</strong></td>
            <td><strong>Perfiles, peldaños y cancelas en ACERO INOXIDABLE</strong> para 2 escalas con crinolina</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">3.700,00 €</td>
            <td class="td-right"><strong style="color: #1e40af;">3.700,00 €</strong></td>
          </tr>
          <tr style="background: #fefce8;">
            <td class="td-center"><strong style="color: #854d0e;">MAT-09</strong></td>
            <td><strong>Materiales y accesorios homologados LUXTOP®</strong> (cable inox 316, postes absorbedores y carros)</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">10.300,00 €</td>
            <td class="td-right"><strong style="color: #854d0e;">10.300,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>MO-01</strong></td>
            <td>Mano de obra especializada montaje HAM obra base (cuadrilla de 4 oficiales ~1050 horas a 28 €/h)</td>
            <td class="td-center">horas</td>
            <td class="td-right">1.050,00</td>
            <td class="td-right">28,00 €</td>
            <td class="td-right"><strong>29.400,00 €</strong></td>
          </tr>
          <tr style="background: #eff6ff;">
            <td class="td-center"><strong style="color: #1e40af;">MO-02</strong></td>
            <td><strong>Mano de obra taller y montaje 2 escalas inox + montaje/ensayo LuxTop</strong> (~100 h)</td>
            <td class="td-center">horas</td>
            <td class="td-right">100,00</td>
            <td class="td-right">28,00 €</td>
            <td class="td-right"><strong style="color: #1e40af;">2.800,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>LOG-01</strong></td>
            <td>Fletes marítimos Península-Menorca para transporte de trailers especiales paneles gran formato</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">5.600,00 €</td>
            <td class="td-right"><strong>5.600,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>LOG-02</strong></td>
            <td>Camión grúa pluma autopropulsada para descarga e izado distribuido sobre pórticos en aeropuerto</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">4.600,00 €</td>
            <td class="td-right"><strong>4.600,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>LOG-03</strong></td>
            <td>Alquiler de plataformas elevadoras móviles de personal (PEMP) articuladas para montaje</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">4.800,00 €</td>
            <td class="td-right"><strong>4.800,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>PRL-01</strong></td>
            <td>Protecciones colectivas: redes anticaídas bajo faldones UNE-EN 1263-1 y líneas de vida EN 795</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">3.800,00 €</td>
            <td class="td-right"><strong>3.800,00 €</strong></td>
          </tr>
          <tr>
            <td class="td-center"><strong>PRL-02</strong></td>
            <td>Tasas de centro de reciclaje y canon vertedero insular de Menorca</td>
            <td class="td-center">lote</td>
            <td class="td-right">1,00</td>
            <td class="td-right">2.400,00 €</td>
            <td class="td-right"><strong>2.400,00 €</strong></td>
          </tr>
          <tr style="background: #f1f5f9; font-weight: 700;">
            <td colspan="5" style="text-align: right; text-transform: uppercase;">COSTE DIRECTO TOTAL GLOBAL (CDT HAM):</td>
            <td class="td-right" style="color: #1e1b4b; font-size: 10px;">{data['escandalloCostesInternosHAM']['costeDirectoTotalGlobal']:,.2f} €</td>
          </tr>
        </tbody>
      </table>

      <div class="margin-card">
        <div class="margin-grid">
          <div>
            <div class="margin-label">Venta Base (BI)</div>
            <div class="margin-val" style="color: #0f172a;">{data['totalesEconomicos']['baseImponibleBase']:,.2f} €</div>
          </div>
          <div>
            <div class="margin-label">Venta Opcionales</div>
            <div class="margin-val" style="color: #1e3a8a;">+{data['totalesEconomicos']['baseImponibleOpcionales']:,.2f} €</div>
          </div>
          <div>
            <div class="margin-label">Coste Directo Total</div>
            <div class="margin-val" style="color: #dc2626;">{data['escandalloCostesInternosHAM']['costeDirectoTotalGlobal']:,.2f} €</div>
          </div>
          <div>
            <div class="margin-label">Margen Bruto HAM</div>
            <div class="margin-val">{data['escandalloCostesInternosHAM']['margenBrutoGlobal']:,.2f} € ({data['escandalloCostesInternosHAM']['margenBrutoPorcentajeVentaGlobal']:.1f}%)</div>
          </div>
        </div>
      </div>
    </div>

    <footer>
      <span>HAM Estructuras Metálicas · Control de Gestión y Costes</span>
      <span>Dossier Técnico Interno Confidencial · P-2026-0044 (REV5) · Página 2 de 2</span>
    </footer>
  </div>
</div>
</body>
</html>
"""

    with open(os.path.join(dir_path, "presupuesto-interno-tecnico.html"), "w", encoding="utf-8") as f:
        f.write(html_interno)
    print("presupuesto-interno-tecnico.html actualizado con éxito.")

def generate_pdfs():
    dir_path = os.path.dirname(__file__)
    
    html_cliente = os.path.join(dir_path, "presupuesto-cliente.html")
    pdf_cliente = os.path.join(dir_path, "presupuesto-cliente.pdf")

    html_interno = os.path.join(dir_path, "presupuesto-interno-tecnico.html")
    pdf_interno = os.path.join(dir_path, "presupuesto-interno-tecnico.pdf")
    
    with sync_playwright() as p:
        browser = p.chromium.launch()

        # PDF Cliente
        page1 = browser.new_page()
        page1.goto(f"file:///{html_cliente.replace(os.sep, '/')}")
        page1.pdf(
            path=pdf_cliente,
            format="A4",
            print_background=True,
            margin={"top": "0mm", "bottom": "0mm", "left": "0mm", "right": "0mm"}
        )
        page1.close()
        print(f"PDF Cliente generado exitosamente en: {pdf_cliente}")

        # PDF Interno
        page2 = browser.new_page()
        page2.goto(f"file:///{html_interno.replace(os.sep, '/')}")
        page2.pdf(
            path=pdf_interno,
            format="A4",
            print_background=True,
            margin={"top": "0mm", "bottom": "0mm", "left": "0mm", "right": "0mm"}
        )
        page2.close()
        print(f"PDF Interno generado exitosamente en: {pdf_interno}")

        browser.close()

if __name__ == "__main__":
    create_excel()
    generate_html_files()
    generate_pdfs()
