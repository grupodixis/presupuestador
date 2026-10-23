import sys
import subprocess
import os

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import openpyxl
except ImportError:
    install('openpyxl')
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

wb = Workbook()
ws = wb.active
ws.title = "Presupuesto Demolicion"

headers = ['Partida', 'Descripción', 'Unidad', 'Cantidad', 'Precio Unitario (€)', 'Total (€)']
ws.append(headers)

# Styling headers
header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
header_font = Font(bold=True)
for col in range(1, 7):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

data = [
    ['01', 'Trabajos previos, limpieza de maleza base, vallado perimetral y señalización (radio 40m)', 'Ud', 1, 650.00],
    ['02', 'Amarre de cable de acero estructural en parte superior y conexión a maquinaria de tiro', 'Ud', 1, 350.00],
    ['03', 'Cortes direccionales (boca de pato) y debilitamiento de patas traseras con oxicorte', 'Ud', 1, 950.00],
    ['04', 'Maniobra de tiro controlado con cable de acero y maquinaria pesada (tractor/cabrestante)', 'Ud', 1, 1400.00],
    ['05', 'Troceado y desguace del depósito y patas a cota cero', 't', 15, 180.00],
    ['06', 'Carga y transporte de chatarra metálica a planta recicladora', 't', 15, 45.00],
    ['07', 'Recuperación por chatarra metálica (Valor estimado a descontar)', 't', 15, -150.00]
]

for row_data in data:
    # Calculate total
    total = row_data[3] * row_data[4]
    row_data.append(total)
    ws.append(row_data)

# Adjust column widths
ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 85
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 20
ws.column_dimensions['F'].width = 15

# Format currency columns
for row in range(2, len(data) + 2):
    ws.cell(row=row, column=5).number_format = '#,##0.00 €'
    ws.cell(row=row, column=6).number_format = '#,##0.00 €'

# Total row
max_row = len(data) + 1
ws.cell(row=max_row+2, column=5, value="TOTAL PRESUPUESTO (€)").font = Font(bold=True)
total_cell = ws.cell(row=max_row+2, column=6, value=f"=SUM(F2:F{max_row})")
total_cell.font = Font(bold=True)
total_cell.number_format = '#,##0.00 €'

# Ensure directory exists
output_dir = r"C:\Users\pablo\Documents\Presupuestador"
os.makedirs(output_dir, exist_ok=True)
file_path = os.path.join(output_dir, "Presupuesto_Demolicion_Tanque.xlsx")

wb.save(file_path)
print(f"Successfully saved to {file_path}")
