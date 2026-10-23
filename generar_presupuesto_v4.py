import sys
import subprocess
import os

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import openpyxl
except ImportError:
    install('openpyxl')
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

wb = Workbook()
ws = wb.active
ws.title = "Presupuesto_Montana"

headers = ['Código', 'Descripción de la Unidad de Obra', 'Unidad', 'Cantidad', 'Precio Unitario (€)', 'Total (€)']
ws.append(headers)

header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
header_font = Font(bold=True)
for col in range(1, 7):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

data = [
    ['01.01', 'Estudio, Plan de Seguridad y Salud, y Coordinación (CSS)', 'Ud', 1, 2540.50],
    ['01.02', 'Implantación de medios preventivos (PRIF, vallado y rescate)', 'Ud', 1, 3615.20],
    ['02.01', 'Adecuación de accesos y movilización de maquinaria (excavadora cadenas)', 'Ud', 1, 6826.00],
    ['02.02', 'Amarre de línea de tracción en cota +25m (Técnicos verticales)', 'Ud', 1, 4150.00],
    ['03.01', 'Apuntalamiento provisional, cortes direccionales y oxicorte base', 'Ud', 1, 3960.00],
    ['03.02', 'Ejecución de maniobra de colapso con tracción (Excavadora CAT 320)', 'Ud', 1, 5745.50],
    ['04.01', 'Reducción volumétrica (troceado manual in situ)', 't', 15, 460.00],
    ['04.02', 'Extracción en dúmperes 4x4 y transporte final a centro autorizado', 't', 15, 371.80]
]

for row_data in data:
    total = row_data[3] * row_data[4]
    row_data.append(total)
    ws.append(row_data)

ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 95
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 20
ws.column_dimensions['F'].width = 15

for row in range(2, len(data) + 2):
    ws.cell(row=row, column=5).number_format = '#,##0.00 €'
    ws.cell(row=row, column=6).number_format = '#,##0.00 €'

max_row = len(data) + 1
ws.cell(row=max_row+2, column=5, value="TOTAL PRESUPUESTO (€)").font = Font(bold=True)
total_cell = ws.cell(row=max_row+2, column=6, value=f"=SUM(F2:F{max_row})")
total_cell.font = Font(bold=True)
total_cell.number_format = '#,##0.00 €'

output_dir = r"C:\Users\pablo\Documents\Presupuestador"
file_path = os.path.join(output_dir, "Presupuesto_Demolicion_Montana.xlsx")

wb.save(file_path)
print(f"Saved to {file_path}")
