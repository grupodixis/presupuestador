import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import re

def main():
    raw_file = r"C:\Users\pablo\Documents\Presupuestador\raw_oyrsa.txt"
    excel_file = r"C:\Users\pablo\Documents\Presupuestador\Presupuesto_OYRSA_Stavion.xlsx"
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presupuesto OYRSA - Stavion"
    
    headers = ["Código", "Concepto", "Ud.", "Cantidad", "Precio Unitario (€)", "Importe (€)"]
    ws.append(["PRESUPUESTO OYRSA - STAVION (Electricidad, Clima y Fontanería)"])
    ws.append([])
    ws.append(headers)
    
    # Style headers
    for col_num, value in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    
    if os.path.exists(raw_file):
        with open(raw_file, "r", encoding="utf-8") as f:
            lines = f.readlines()
        
        # Regex to parse lines: Code Partida Unit Description Quantity 0,00...
        pattern = re.compile(r"^(\S+)\s+Partida\s+(\S+)\s+(.*?)\s+([0-9\.]+,\d{2}|\d+,\d{2})\s+0,00\s+0,00")
        
        row_idx = 4
        for line in lines:
            line = line.strip()
            if not line:
                continue
            match = pattern.match(line)
            if match:
                codigo = match.group(1)
                unidad = match.group(2)
                descripcion = match.group(3).strip()
                cantidad_str = match.group(4).replace('.', '').replace(',', '.')
                cantidad = float(cantidad_str)
                
                # Formula for importe = Cantidad * Precio Unitario
                formula = f"=D{row_idx}*E{row_idx}"
                ws.append([codigo, descripcion, unidad, cantidad, 0.0, formula])
                row_idx += 1
                
        # Subtotal, IVA, Total
        ws.append([])
        subtotal_row = row_idx + 1
        ws.append(["", "", "", "", "SUBTOTAL", f"=SUM(F4:F{row_idx-1})"])
        ws.append(["", "", "", "", "IVA (21%)", f"=F{subtotal_row}*0.21"])
        ws.append(["", "", "", "", "TOTAL PRESUPUESTO", f"=F{subtotal_row}+F{subtotal_row+1}"])
        
        # Style Totals
        for r in range(subtotal_row, subtotal_row + 3):
            ws.cell(row=r, column=5).font = Font(bold=True)
            ws.cell(row=r, column=6).font = Font(bold=True)
            ws.cell(row=r, column=6).number_format = '#,##0.00 €'
        
        # Formatting
        for r in range(4, row_idx):
            ws.cell(row=r, column=4).number_format = '#,##0.00'
            ws.cell(row=r, column=5).number_format = '#,##0.00 €'
            ws.cell(row=r, column=6).number_format = '#,##0.00 €'
            
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 80)
            ws.column_dimensions[column].width = adjusted_width
            
        wb.save(excel_file)
        print(f"Archivo guardado: {excel_file}")
    else:
        print("Error: No se encontro raw_oyrsa.txt")

if __name__ == '__main__':
    main()
