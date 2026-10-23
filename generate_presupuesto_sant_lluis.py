import subprocess
import sys
import os

def install_and_import(package):
    try:
        __import__(package)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])

install_and_import('openpyxl')

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Presupuesto Sant Lluís"

# Configuración de bordes y rellenos
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
client_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

# Información del cliente
ws.merge_cells('A1:F1')
ws['A1'] = "PRESUPUESTO: CONJUNTO DE CARPINTERÍA (PE2 + PE3 UNIFICADAS)"
ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
ws['A1'].fill = header_fill
ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

ws['A3'] = "Cliente:"
ws['B3'] = "Sonia Arjona"
ws['A4'] = "Ubicación:"
ws['B4'] = "Sant Lluís, Menorca (Cambio de uso de garaje a vivienda)"

for r in range(3, 5):
    ws[f'A{r}'].font = Font(bold=True)
    for c in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws[f'{c}{r}'].fill = client_fill

# Headers del presupuesto
headers = ["Concepto", "Descripción", "Medidas / Cantidad", "Uds", "Precio Unitario", "Total"]
ws.append([]) # Fila vacía (fila 5)
ws.append(headers) # Fila 6

for col_num, value in enumerate(headers, 1):
    cell = ws.cell(row=6, column=col_num)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

# Datos
data = [
    [
        "Estructura Principal",
        "Estructura de aluminio con Rotura de Puente Térmico (RPT). Color: Blanco (BLANC). Incluye montante vertical central continuo (suelo a techo) para fijación posterior de tabiquería seca de separación interna.",
        "Hueco: 3,50 m (ancho) x 3,02 m (alto)",
        "1",
        "",
        ""
    ],
    [
        "Zona Peatonal (PE3)",
        "Puerta de 1 hoja batiente. Panel ciego de aluminio RPT blanco. Transmitancia térmica U ≤ 3.2 W/m²K.",
        "85 cm (ancho) x 211 cm (alto)",
        "1",
        "",
        ""
    ],
    [
        "Zona Garaje (PE2)",
        "Puerta seccional motorizada. Incluye motor de techo y 2 mandos a distancia. Paneles en color blanco a juego con la puerta peatonal. Ventilación obligatoria mediante lamas troqueladas (superficie mínima 480 cm²). Transmitancia térmica U ≤ 2.2 W/m²K.",
        "254 cm (ancho) x 211 cm (alto)",
        "1",
        "",
        ""
    ],
    [
        "Fijos Superiores Acristalados",
        "Dos paños fijos independientes (divididos por montante central). Vidrio doble transparente con propiedades de Baja Emisividad (BAIX EMISSIU) / Low-E.",
        "Alturas: 91 cm aprox.",
        "2",
        "",
        ""
    ],
    [
        "Instalación",
        "Medición definitiva en obra, transporte e instalación completa del conjunto de carpintería.",
        "Servicio",
        "1",
        "",
        ""
    ]
]

for row_idx, row_data in enumerate(data, start=7):
    ws.append(row_data)
    for col_idx in range(1, 7):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.border = thin_border
        if col_idx == 2:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        elif col_idx in [3, 4, 5, 6]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(vertical="center")

# Anchos de columna
column_widths = {
    'A': 25,
    'B': 70,
    'C': 25,
    'D': 8,
    'E': 15,
    'F': 15
}

for col, width in column_widths.items():
    ws.column_dimensions[col].width = width

# Total Final (filas abajo)
last_row = ws.max_row
ws.merge_cells(f'A{last_row+2}:E{last_row+2}')
ws[f'A{last_row+2}'] = "BASE IMPONIBLE:"
ws[f'A{last_row+2}'].font = Font(bold=True)
ws[f'A{last_row+2}'].alignment = Alignment(horizontal="right")
ws[f'F{last_row+2}'].border = thin_border

ws.merge_cells(f'A{last_row+3}:E{last_row+3}')
ws[f'A{last_row+3}'] = "IVA (21%):"
ws[f'A{last_row+3}'].font = Font(bold=True)
ws[f'A{last_row+3}'].alignment = Alignment(horizontal="right")
ws[f'F{last_row+3}'].border = thin_border

ws.merge_cells(f'A{last_row+4}:E{last_row+4}')
ws[f'A{last_row+4}'] = "TOTAL PRESUPUESTO:"
ws[f'A{last_row+4}'].font = Font(bold=True)
ws[f'A{last_row+4}'].alignment = Alignment(horizontal="right")
ws[f'F{last_row+4}'].border = thin_border
ws[f'F{last_row+4}'].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

excel_path = r"C:\Users\pablo\Documents\Presupuestador\Presupuesto_Carpinteria_Sant_Lluis.xlsx"
wb.save(excel_path)
print(f"Archivo Excel creado exitosamente en {excel_path}")
