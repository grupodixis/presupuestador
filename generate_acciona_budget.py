import subprocess
import sys

def install_and_import(package):
    try:
        __import__(package)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])

install_and_import('openpyxl')

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Presupuesto Acciona Aguas"

# Headers
headers = ["Opción", "Concepto", "Ud.", "Cantidad", "Precio Unitario (€)", "Importe (€)"]
ws.append(["PRESUPUESTO PARA ACCIONA AGUAS SA"])
ws.append(["Instalación de 4 puertas seccionales (con y sin opción motorizada)"])
ws.append([])
ws.append(headers)

for col_num, value in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_num)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    cell.alignment = Alignment(horizontal="center")

ws.cell(row=1, column=1).font = Font(bold=True, size=14)
ws.cell(row=2, column=1).font = Font(bold=True, size=12)

# Data
data = [
    ["OPCIÓN A (SIN MOTOR)", "Suministro Puerta Seccional Ind. EUROLINES con paso peatonal", "Ud.", 2, 2328.90, 4657.80],
    ["OPCIÓN A (SIN MOTOR)", "Suministro Puerta Seccional Ind. EUROLINES sin paso peatonal", "Ud.", 2, 1369.40, 2738.80],
    ["OPCIÓN A (SIN MOTOR)", "Mano de obra instalación mecánica", "H.", 64, 61.68, 3947.52],
    ["", "Subtotal Ejecución Material (Opción A)", "", "", "", 11344.12],
    ["", "Gastos generales (13%)", "", "", "", 1474.74],
    ["", "Beneficio industrial (6%)", "", "", "", 680.65],
    ["", "TOTAL OPCIÓN A (Sin IVA)", "", "", "", 13499.51],
    ["", "", "", "", "", ""],
    ["OPCIÓN B (CON MOTOR)", "Kit Motorización puertas con paso peatonal (motor, fotocélulas, emisor)", "Ud.", 2, 1368.00, 2736.00],
    ["OPCIÓN B (CON MOTOR)", "Kit Motorización puertas sin paso peatonal (motor, fotocélulas, emisor)", "Ud.", 2, 1305.00, 2610.00],
    ["OPCIÓN B (CON MOTOR)", "Material pequeña instalación eléctrica para motores (cableado, tubo, protecciones)", "Ud.", 4, 150.00, 600.00],
    ["OPCIÓN B (CON MOTOR)", "Mano de obra instalación eléctrica y automatismos", "H.", 32, 24.27, 776.64],
    ["", "Subtotal Ejecución Material (Motorización)", "", "", "", 6722.64],
    ["", "Gastos generales (13%)", "", "", "", 873.94],
    ["", "Beneficio industrial (6%)", "", "", "", 403.36],
    ["", "TOTAL OPCIÓN B (Opcional - Sin IVA)", "", "", "", 7999.94],
    ["", "", "", "", "", ""],
    ["", "TOTAL OPCIÓN A + B (Sin IVA)", "", "", "", 21499.45],
    ["", "IVA (21%)", "", "", "", 4514.88],
    ["", "TOTAL PRESUPUESTO (OPCIÓN A + B) CON IVA", "", "", "", 26014.33]
]

for row in data:
    ws.append(row)

# Styling
for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=5, max_col=6):
    for cell in row:
        if isinstance(cell.value, (int, float)):
            cell.number_format = '#,##0.00 €'

for col in ws.columns:
    max_length = 0
    column = col[0].column_letter # Get the column name
    for cell in col:
        try: # Necessary to avoid error on empty cells
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = (max_length + 2)
    ws.column_dimensions[column].width = adjusted_width

# Totals styling
for row_num in range(5, ws.max_row + 1):
    if ws.cell(row=row_num, column=2).value and "TOTAL" in str(ws.cell(row=row_num, column=2).value):
        ws.cell(row=row_num, column=2).font = Font(bold=True)
        if isinstance(ws.cell(row=row_num, column=6).value, (int, float)):
            ws.cell(row=row_num, column=6).font = Font(bold=True)
    if ws.cell(row=row_num, column=2).value and "Subtotal" in str(ws.cell(row=row_num, column=2).value):
        ws.cell(row=row_num, column=2).font = Font(bold=True, italic=True)
        if isinstance(ws.cell(row=row_num, column=6).value, (int, float)):
            ws.cell(row=row_num, column=6).font = Font(bold=True, italic=True)

wb.save(r"C:\Users\pablo\Documents\Presupuestador\Presupuesto_Acciona_Aguas.xlsx")
print("Archivo Excel creado exitosamente.")
