import sys
import subprocess
import os

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import openpyxl
except ImportError:
    install('openpyxl')
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

wb = Workbook()
ws = wb.active
ws.title = "Presupuesto_Montana"

headers = ['Partida', 'Descripción', 'Unidad', 'Cantidad', 'Precio Unitario (€)', 'Total (€)']
ws.append(headers)

header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
header_font = Font(bold=True)
for col in range(1, 7):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

data = [
    ['01', 'Estudio, Plan de Seguridad y Salud, y Coordinación de SS', 'Ud', 1, 2500.00],
    ['02', 'Implantación SS: Prevención incendios, vallado perimetral y medios de rescate', 'Ud', 1, 3500.00],
    ['03', 'Adecuación de accesos forestales y movilización de maquinaria especial (orugas)', 'Ud', 1, 6500.00],
    ['04', 'Amarre de línea de tiro en altura mediante técnicos de trabajos verticales (IRATA)', 'Ud', 1, 4000.00],
    ['05', 'Cortes direccionales y debilitamiento de soportes in situ con oxicorte', 'Ud', 1, 3800.00],
    ['06', 'Maniobra de tiro controlado con cabrestante y maquinaria forestal pesada', 'Ud', 1, 5500.00],
    ['07', 'Troceado exhaustivo manual in situ con oxicorte para posibilitar carga ligera', 't', 15, 600.00],
    ['08', 'Extracción por pista con dúmperes 4x4 y transporte a gestor autorizado', 't', 15, 350.00],
    ['09', 'Recuperación por chatarra metálica (Valor de mercado ajustado por logística forestal)', 't', 15, -70.00]
]

for row_data in data:
    total = row_data[3] * row_data[4]
    row_data.append(total)
    ws.append(row_data)

ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 95
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 20
ws.column_dimensions['F'].width = 15

for row in range(2, len(data) + 2):
    ws.cell(row=row, column=5).number_format = '#,##0.00 €'
    ws.cell(row=row, column=6).number_format = '#,##0.00 €'

max_row = len(data) + 1
ws.cell(row=max_row+2, column=5, value="TOTAL PRESUPUESTO (€)").font = Font(bold=True)
total_cell = ws.cell(row=max_row+2, column=6, value=f"=SUM(F2:F{max_row})")
total_cell.font = Font(bold=True)
total_cell.number_format = '#,##0.00 €'

output_dir = r"C:\Users\pablo\Documents\Presupuestador"
file_path = os.path.join(output_dir, "Presupuesto_Demolicion_Montana.xlsx")

wb.save(file_path)
print(f"Saved to {file_path}")
