from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path


ROOT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("correos")
LIMIT = int(sys.argv[2]) if len(sys.argv) > 2 and sys.argv[2].isdigit() else 0
ROOT = ROOT.resolve()
PROCESADO = ROOT / "_procesado"
ATTACHMENTS = PROCESADO / "adjuntos-candidatos"
TEXTS = PROCESADO / "textos-adjuntos"

PRICE_LINE_KEYWORDS = [
    "aluminio",
    "acero",
    "inox",
    "aisi",
    "chapa",
    "perfil",
    "tubo",
    "pletina",
    "lama",
    "mallorquina",
    "menorquina",
    "barandilla",
    "puerta",
    "estructura",
    "sandwich",
    "sándwich",
    "panel",
    "cerramiento",
    "viga",
    "ipe",
    "hea",
    "heb",
    "mano de obra",
    "montaje",
    "transporte",
    "lacado",
    "galvanizado",
    "tornilleria",
    "tornillería",
    "cristal",
    "vidrio",
    "prfv",
    "pvc",
]

NOISE_KEYWORDS = [
    "base imponible",
    "cuota iva",
    "total euros",
    "total bruto",
    "total general",
    "reg.mercantil",
    "proteccion de datos",
    "protección de datos",
    "iban",
    "vencimiento",
]

MONEY_RE = re.compile(r"(?<![\w/])(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d{2}|\.\d{2})(?![\w/])")
DATE_RE = re.compile(r"\b(?:\d{1,2}[/-]\d{1,2}[/-](?:20)?\d{2}|20\d{2}[/-]\d{1,2}[/-]\d{1,2})\b")


@dataclass
class AttachmentMeta:
    source_file: str
    email_date: str = ""
    sender: str = ""
    subject: str = ""
    score: str = ""
    reasons: str = ""


def safe_rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def parse_number(value: str) -> float | None:
    value = value.strip()
    if "," in value:
        value = value.replace(".", "").replace(",", ".")
    try:
        return float(value)
    except ValueError:
        return None


def infer_year(date_text: str, email_date: str) -> str:
    if date_text:
        year = re.search(r"20\d{2}", date_text)
        if year:
            return year.group(0)
    if email_date:
        return email_date[:4]
    return ""


def confidence_from_dates(doc_year: str, email_year: str) -> str:
    year = doc_year or email_year
    if not year:
        return "pendiente"
    try:
        numeric = int(year)
    except ValueError:
        return "pendiente"
    if numeric >= 2026:
        return "pendiente"
    return "antiguo"


def load_metadata() -> dict[str, AttachmentMeta]:
    meta: dict[str, AttachmentMeta] = {}
    csv_path = PROCESADO / "candidatos-precios.csv"
    if not csv_path.exists():
        return meta
    with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            saved = row.get("saved_attachments", "")
            for item in saved.split(" | "):
                if not item.strip():
                    continue
                key = str(Path(item.strip()).resolve()).lower()
                meta[key] = AttachmentMeta(
                    source_file=item.strip(),
                    email_date=row.get("email_date", ""),
                    sender=row.get("from", ""),
                    subject=row.get("subject", ""),
                    score=row.get("score", ""),
                    reasons=row.get("reasons", ""),
                )
    return meta


def extract_pdf(path: Path) -> str:
    try:
        import pypdf

        reader = pypdf.PdfReader(str(path))
        pages = []
        for page in reader.pages[:3]:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    except Exception as exc:
        return f"[ERROR PDF {type(exc).__name__}: {exc}]"


def extract_docx(path: Path) -> str:
    try:
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables[:10]:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        return f"[ERROR DOCX {type(exc).__name__}: {exc}]"


def extract_xlsx(path: Path) -> str:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        lines = []
        for ws in wb.worksheets[:6]:
            lines.append(f"# SHEET {ws.title}")
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 250), values_only=True):
                values = [normalize_text(str(v)) if v is not None else "" for v in row]
                if any(values):
                    lines.append(" | ".join(values))
        return "\n".join(lines)
    except Exception as exc:
        return f"[ERROR XLSX {type(exc).__name__}: {exc}]"


def extract_zip(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
        return "\n".join(f"ZIP: {name}" for name in names[:200])
    except Exception as exc:
        return f"[ERROR ZIP {type(exc).__name__}: {exc}]"


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return extract_xlsx(path)
    if suffix == ".zip":
        return extract_zip(path)
    return ""


def text_path_for(path: Path) -> Path:
    digest = hashlib.sha1(str(path).encode("utf-8", errors="ignore")).hexdigest()[:10]
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", path.stem)[:80]
    return TEXTS / f"{stem}-{digest}.txt"


def line_score(line: str) -> int:
    lower = line.lower()
    score = 0
    score += len(MONEY_RE.findall(line)) * 3
    for keyword in PRICE_LINE_KEYWORDS:
        if keyword in lower:
            score += 2
    for keyword in NOISE_KEYWORDS:
        if keyword in lower:
            score -= 4
    if len(line) < 20:
        score -= 3
    if len(line) > 500:
        score -= 2
    return score


def guess_area(text: str) -> str:
    lower = text.lower()
    if any(k in lower for k in ["aluminio", "mallorquina", "menorquina", "cortizo", "pvc"]):
        return "aluminio"
    if any(k in lower for k in ["fontaneria", "fontanería", "saneamiento", "bombeo"]):
        return "fontaneria"
    if any(k in lower for k in ["clima", "ventilacion", "ventilación", "frigor"]):
        return "clima"
    if any(k in lower for k in ["electric", "iluminacion", "iluminación", "cuadro"]):
        return "instalaciones_electricas"
    if any(k in lower for k in ["acero", "inox", "barandilla", "estructura", "metal", "hierro"]):
        return "carpinteria_metalica"
    return "pendiente"


def guess_unit(line: str) -> str:
    lower = line.lower()
    for unit in ["kg", "ml", "m2", "m²", "ud", "uds", "h", "m3", "m³", "lote"]:
        if re.search(rf"\b{re.escape(unit)}\b", lower):
            return unit.replace("m²", "m2").replace("m³", "m3")
    return ""


def candidate_lines(text: str, path: Path, meta: AttachmentMeta) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    doc_date_match = DATE_RE.search(text[:5000])
    doc_date = doc_date_match.group(0) if doc_date_match else ""
    doc_year = infer_year(doc_date, meta.email_date)
    confidence = confidence_from_dates(doc_year, meta.email_date[:4])

    for raw in text.splitlines():
        line = normalize_text(raw)
        if not line:
            continue
        monies = MONEY_RE.findall(line)
        score = line_score(line)
        if score < 8 or len(monies) < 1:
            continue
        values = [parse_number(m) for m in monies]
        values = [v for v in values if v is not None]
        if not values:
            continue
        unit_price = values[-2] if len(values) >= 2 else ""
        total = values[-1]
        quantity = values[-3] if len(values) >= 3 else ""
        description = line
        for m in monies[-3:]:
            description = description.replace(m, " ")
        description = normalize_text(description)
        rows.append(
            {
                "line_score": str(score),
                "email_date": meta.email_date,
                "document_date_hint": doc_date,
                "document_year": doc_year,
                "confidence": confidence,
                "area": guess_area(" ".join([line, path.name, meta.subject, meta.sender])),
                "supplier_hint": meta.sender,
                "subject": meta.subject,
                "attachment": str(path),
                "attachment_name": path.name,
                "unit_hint": guess_unit(line),
                "description": description[:400],
                "quantity_hint": str(quantity),
                "unit_price_hint": str(unit_price),
                "total_hint": str(total),
                "raw_line": line[:800],
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict[str, str]], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def is_reviewable_price_row(row: dict[str, str]) -> bool:
    desc = (row.get("description") or "").lower()
    raw = (row.get("raw_line") or "").lower()
    try:
        unit_price = float(row.get("unit_price_hint") or "0")
        total = float(row.get("total_hint") or "0")
    except ValueError:
        return False
    if unit_price <= 0 or total <= 0:
        return False
    if len(desc) < 8:
        return False
    if any(k in desc or k in raw for k in NOISE_KEYWORDS):
        return False
    if re.fullmatch(r"[\d\s,\.€|-]+", desc):
        return False
    return True


def main() -> int:
    TEXTS.mkdir(parents=True, exist_ok=True)
    metadata = load_metadata()
    files = [
        p
        for p in ATTACHMENTS.rglob("*")
        if p.is_file() and p.suffix.lower() in {".pdf", ".docx", ".xlsx", ".xlsm", ".xls", ".zip"}
        and p.stat().st_size <= 20_000_000
    ]
    files.sort(
        key=lambda p: (
            metadata.get(str(p.resolve()).lower(), AttachmentMeta(source_file=str(p))).email_date,
            int(metadata.get(str(p.resolve()).lower(), AttachmentMeta(source_file=str(p))).score or 0),
            str(p),
        ),
        reverse=True,
    )
    if LIMIT:
        files = files[:LIMIT]

    attachment_rows: list[dict[str, str]] = []
    price_rows: list[dict[str, str]] = []

    for idx, path in enumerate(files, 1):
        meta = metadata.get(str(path.resolve()).lower(), AttachmentMeta(source_file=str(path)))
        text = extract_text(path)
        tp = text_path_for(path)
        tp.write_text(text, encoding="utf-8", errors="ignore")
        has_text = bool(text.strip()) and not text.startswith("[ERROR")
        attachment_rows.append(
            {
                "email_date": meta.email_date,
                "from": meta.sender,
                "subject": meta.subject,
                "attachment": str(path),
                "type": path.suffix.lower(),
                "size_bytes": str(path.stat().st_size),
                "text_file": str(tp),
                "has_text": "yes" if has_text else "no",
                "chars": str(len(text)),
                "error": text if text.startswith("[ERROR") else "",
            }
        )
        if has_text:
            price_rows.extend(candidate_lines(text, path, meta))
        if idx % 100 == 0:
            print(f"Procesados {idx}/{len(files)}")

    attachment_fields = [
        "email_date",
        "from",
        "subject",
        "attachment",
        "type",
        "size_bytes",
        "text_file",
        "has_text",
        "chars",
        "error",
    ]
    price_fields = [
        "line_score",
        "email_date",
        "document_date_hint",
        "document_year",
        "confidence",
        "area",
        "supplier_hint",
        "subject",
        "attachment",
        "attachment_name",
        "unit_hint",
        "description",
        "quantity_hint",
        "unit_price_hint",
        "total_hint",
        "raw_line",
    ]
    write_csv(PROCESADO / "adjuntos-extraidos.csv", attachment_rows, attachment_fields)
    write_csv(PROCESADO / "lineas-precio-candidatas.csv", price_rows, price_fields)
    filtered_rows = [row for row in price_rows if is_reviewable_price_row(row)]
    write_csv(PROCESADO / "lineas-precio-filtradas.csv", filtered_rows, price_fields)

    by_area: dict[str, int] = {}
    by_conf: dict[str, int] = {}
    for row in price_rows:
        by_area[row["area"]] = by_area.get(row["area"], 0) + 1
        by_conf[row["confidence"]] = by_conf.get(row["confidence"], 0) + 1

    md = [
        "# Extraccion automatica de lineas de precio",
        "",
        f"- Generado: {datetime.now().isoformat(timespec='seconds')}",
        f"- Adjuntos procesados: {len(attachment_rows)}",
        f"- Lineas candidatas: {len(price_rows)}",
        f"- Lineas filtradas revisables: {len(filtered_rows)}",
        "",
        "## Por area",
        "",
    ]
    for area, count in sorted(by_area.items(), key=lambda item: item[1], reverse=True):
        md.append(f"- {area}: {count}")
    md.extend(["", "## Por confianza", ""])
    for conf, count in sorted(by_conf.items(), key=lambda item: item[1], reverse=True):
        md.append(f"- {conf}: {count}")
    md.extend(["", "## Primeras lineas candidatas", ""])
    for row in sorted(price_rows, key=lambda r: int(r["line_score"]), reverse=True)[:80]:
        md.extend(
            [
                f"### {row['document_year'] or 'sin fecha'} - {row['area']} - {row['attachment_name']}",
                "",
                f"- Confianza: {row['confidence']}",
                f"- Proveedor/remitente: {row['supplier_hint']}",
                f"- Descripcion: {row['description']}",
                f"- Unidad: {row['unit_hint'] or 'pendiente'}",
                f"- Cantidad detectada: {row['quantity_hint'] or 'pendiente'}",
                f"- Precio unitario detectado: {row['unit_price_hint'] or 'pendiente'}",
                f"- Total detectado: {row['total_hint']}",
                f"- Fuente: `{safe_rel(Path(row['attachment']))}`",
                "",
            ]
        )
    (PROCESADO / "lineas-precio-candidatas.md").write_text("\n".join(md), encoding="utf-8")

    print(f"Adjuntos procesados: {len(attachment_rows)}")
    print(f"Lineas candidatas: {len(price_rows)}")
    print(f"Lineas filtradas revisables: {len(filtered_rows)}")
    print(PROCESADO / "adjuntos-extraidos.csv")
    print(PROCESADO / "lineas-precio-candidatas.csv")
    print(PROCESADO / "lineas-precio-filtradas.csv")
    print(PROCESADO / "lineas-precio-candidatas.md")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
